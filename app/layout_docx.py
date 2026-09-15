# -*- coding: utf-8 -*-
"""扫描 PDF → 保留原页版式的 Word：每页原图垫底 + OCR 文本框按 bbox 覆盖。"""
from __future__ import annotations

from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.shared import Emu, Pt

from .scan_text_normalize import collapse_cjk_spaces

_WP = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
_EMU_PER_PT = 12700
_SHAPE_BASE = 1025


def _pt_emu(pt: float) -> int:
    return max(1, int(round(float(pt) * _EMU_PER_PT)))


def _xml_text(s: str) -> str:
    return escape(str(s or ""), {"'": "&apos;", '"': "&quot;"})


def _zero_spacing(paragraph, *, exact_twips: int = 1) -> None:
    """浮动图/文本框所在段不占版心：行距压到 1 twip。"""
    pf = paragraph.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pPr = paragraph._p.get_or_add_pPr()
    spacing = pPr.find(qn("w:spacing"))
    line = str(max(1, int(exact_twips)))
    if spacing is None:
        spacing = parse_xml(
            '<w:spacing xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
            f' w:before="0" w:after="0" w:line="{line}" w:lineRule="exact"/>'
        )
        pPr.append(spacing)
    else:
        spacing.set(qn("w:before"), "0")
        spacing.set(qn("w:after"), "0")
        spacing.set(qn("w:line"), line)
        spacing.set(qn("w:lineRule"), "exact")


def _set_page(section, width_pt: float, height_pt: float) -> None:
    section.page_width = Emu(_pt_emu(width_pt))
    section.page_height = Emu(_pt_emu(height_pt))
    section.left_margin = Emu(0)
    section.right_margin = Emu(0)
    section.top_margin = Emu(0)
    section.bottom_margin = Emu(0)
    section.header_distance = Emu(0)
    section.footer_distance = Emu(0)
    section.gutter = Emu(0)


def _float_picture_behind_page(run, cx: int, cy: int) -> None:
    """把 python-docx 的 inline 图片改成相对页面、衬于文字后的浮动图。"""
    drawing = run._r.find(qn("w:drawing"))
    if drawing is None:
        return
    inline = drawing.find(qn("wp:inline"))
    if inline is None:
        return
    anchor = parse_xml(
        f'<wp:anchor xmlns:wp="{_WP}" xmlns:a="{_A}" distT="0" distB="0" distL="0" distR="0"'
        f' simplePos="0" relativeHeight="0" behindDoc="1" locked="1" layoutInCell="1" allowOverlap="1">'
        f'<wp:simplePos x="0" y="0"/>'
        f'<wp:positionH relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionH>'
        f'<wp:positionV relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionV>'
        f'<wp:extent cx="{cx}" cy="{cy}"/>'
        f'<wp:effectExtent l="0" t="0" r="0" b="0"/>'
        f'<wp:wrapNone/>'
        f"</wp:anchor>"
    )
    for child in list(inline):
        tag = child.tag
        if tag in (qn("wp:extent"), qn("wp:effectExtent")):
            continue
        anchor.append(child)
    drawing.remove(inline)
    drawing.append(anchor)


def _textbox_pict(
    text: str,
    left_pt: float,
    top_pt: float,
    width_pt: float,
    height_pt: float,
    shape_id: int,
    font_pt: float,
    include_shapetype: bool,
) -> object:
    sz = max(16, int(round(font_pt * 2)))  # half-points
    style = (
        f"position:absolute;margin-left:{left_pt:.2f}pt;margin-top:{top_pt:.2f}pt;"
        f"width:{width_pt:.2f}pt;height:{height_pt:.2f}pt;z-index:{shape_id};"
        "mso-position-horizontal-relative:page;mso-position-vertical-relative:page;"
        "v-text-anchor:top"
    )
    shapetype = ""
    if include_shapetype:
        shapetype = (
            '<v:shapetype id="_x0000_t202" coordsize="21600,21600" o:spt="202" '
            'path="m,l,21600r21600,l21600,xe">'
            '<v:stroke joinstyle="miter"/>'
            '<v:path gradientshapeok="t" o:connecttype="rect"/>'
            "</v:shapetype>"
        )
    xml = (
        '<w:pict xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
        ' xmlns:v="urn:schemas-microsoft-com:vml"'
        ' xmlns:o="urn:schemas-microsoft-com:office:office">'
        f"{shapetype}"
        f'<v:shape id="_x0000_s{shape_id}" type="#_x0000_t202" style="{style}"'
        ' filled="t" fillcolor="white" stroked="f" o:allowincell="f">'
        '<v:fill on="t" color="white" opacity="1"/>'
        '<v:textbox inset="0.6pt,0.2pt,0.6pt,0.2pt">'
        "<w:txbxContent>"
        "<w:p>"
        '<w:pPr><w:spacing w:before="0" w:after="0" w:line="240" w:lineRule="auto"/>'
        '<w:ind w:left="0" w:right="0"/></w:pPr>'
        "<w:r>"
        "<w:rPr>"
        f'<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:eastAsia="宋体"/>'
        f'<w:sz w:val="{sz}"/><w:szCs w:val="{sz}"/>'
        "</w:rPr>"
        f'<w:t xml:space="preserve">{_xml_text(text)}</w:t>'
        "</w:r>"
        "</w:p>"
        "</w:txbxContent>"
        "</v:textbox>"
        "</v:shape>"
        "</w:pict>"
    )
    return parse_xml(xml)


def _map_box(el: dict, src_w: float, src_h: float, page_w: float, page_h: float) -> tuple[float, float, float, float]:
    x0, y0, x1, y1 = el["x0"], el["y0"], el["x1"], el["y1"]
    left = x0 / src_w * page_w
    top = y0 / src_h * page_h
    width = (x1 - x0) / src_w * page_w
    height = (y1 - y0) / src_h * page_h
    pad = 1.1
    left = max(0.0, left - pad)
    top = max(0.0, top - pad * 0.6)
    width = min(page_w - left, width + pad * 2)
    height = min(page_h - top, height + pad * 1.2)
    return left, top, max(6.0, width), max(7.0, height)


def _keep_element(el: dict, src_w: float, src_h: float) -> bool:
    text = collapse_cjk_spaces(str(el.get("text") or "")).strip()
    if not text:
        return False
    w = float(el["x1"]) - float(el["x0"])
    h = float(el["y1"]) - float(el["y0"])
    if w < 4 or h < 4:
        return False
    if src_w > 0 and src_h > 0 and (w * h) / (src_w * src_h) > 0.82:
        return False
    return True


def build_layout_preserving_docx(
    dst: Path,
    page_pngs: list[tuple[int, bytes]],
    page_sizes_pt: dict[int, tuple[float, float]],
    layout_pages: dict[int, dict] | None = None,
) -> dict:
    """写出「原页图 + 定位文本框」的 Word。返回 {ok, pages, textboxes}。"""
    dst = Path(dst)
    items = [(int(pg), png) for pg, png in page_pngs if png]
    if not items:
        return {"ok": False, "pages": 0, "textboxes": 0, "error": "没有页面图像"}
    layout_pages = layout_pages or {}
    dst.parent.mkdir(parents=True, exist_ok=True)
    if dst.exists():
        dst.unlink()

    doc = Document()
    sect0 = doc.sections[0]
    body = doc.element.body
    first_p = body.find(qn("w:p"))
    if first_p is not None:
        body.remove(first_p)

    textboxes = 0
    shape_id = _SHAPE_BASE
    first_box = True

    for idx, (pg, png) in enumerate(items):
        w_pt, h_pt = page_sizes_pt.get(pg) or (595.3, 841.9)
        w_pt = float(w_pt or 595.3)
        h_pt = float(h_pt or 841.9)
        if idx == 0:
            section = sect0
        else:
            section = doc.add_section(WD_SECTION.NEW_PAGE)
            # add_section 在上一节末尾插入空段承载 sectPr，压成零高以免多出空白页
            for prev in reversed(list(doc.paragraphs)):
                pPr = prev._p.find(qn("w:pPr"))
                if pPr is not None and pPr.find(qn("w:sectPr")) is not None:
                    _zero_spacing(prev)
                    break
        _set_page(section, w_pt, h_pt)

        bg = doc.add_paragraph()
        _zero_spacing(bg)
        bg.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = bg.add_run()
        cx, cy = _pt_emu(w_pt), _pt_emu(h_pt)
        run.add_picture(BytesIO(png), width=Emu(cx), height=Emu(cy))
        _float_picture_behind_page(run, cx, cy)

        rec = layout_pages.get(pg) or {}
        src_w = float(rec.get("width") or 0) or w_pt
        src_h = float(rec.get("height") or 0) or h_pt
        for el in rec.get("elements") or []:
            if not _keep_element(el, src_w, src_h):
                continue
            text = collapse_cjk_spaces(str(el.get("text") or "")).strip()
            left, top, width, height = _map_box(el, src_w, src_h, w_pt, h_pt)
            font_pt = min(16.0, max(8.0, height * 0.72))
            if height > 28:
                font_pt = min(font_pt, 11.0)
            shape_id += 1
            box_p = doc.add_paragraph()
            _zero_spacing(box_p)
            box_run = box_p.add_run()
            box_run._r.append(
                _textbox_pict(text, left, top, width, height, shape_id, font_pt, first_box)
            )
            first_box = False
            textboxes += 1

    # 去掉 python-docx 默认 sectPr 可能残留的页边距
    _set_page(doc.sections[-1], *page_sizes_pt.get(items[-1][0], (595.3, 841.9)))
    doc.save(str(dst))
    return {"ok": True, "pages": len(items), "textboxes": textboxes}
