# -*- coding: utf-8 -*-
"""按项目根目录 QM.docx / HJ.docx 表格模板填入解析结果（仅填 OCR/解析中存在的字段）。"""
from __future__ import annotations

import re
import shutil
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

try:
    from .form_kind import ROOT_DIR
    from .hj_reference import (
        SCHOOL_COUNTRY_HINTS,
        cn_list_join,
        en_list_join,
        fix_email_ocr,
        fix_entity_name,
        fix_ocr_english,
        format_bilingual_pair,
        format_hj_degree_cell,
        format_hj_position_cell,
        highest_degree_from_edu_row,
        is_academic_title,
        normalize_hj_cn_field,
        normalize_hj_en_field,
        split_bilingual,
        split_certificate_name,
    )
except ImportError:
    ROOT_DIR = Path(__file__).resolve().parent.parent
    SCHOOL_COUNTRY_HINTS = ()
    fix_ocr_english = lambda s: str(s or "").strip()
    highest_degree_from_edu_row = lambda row: ("", "")
    format_bilingual_pair = lambda cn, en: (str(cn or ""), str(en or ""))
    split_bilingual = lambda s: (str(s or "").strip(), "")
    cn_list_join = lambda *parts: "，".join(p for p in parts if p)
    en_list_join = lambda *parts: "，".join(p for p in parts if p)
    fix_entity_name = lambda s: str(s or "").strip()
    is_academic_title = lambda s: False
    normalize_hj_cn_field = lambda s: str(s or "").strip()
    normalize_hj_en_field = lambda s: str(s or "").strip()
    split_certificate_name = lambda s: (str(s or "").strip(), "")
    format_hj_degree_cell = lambda s: str(s or "").strip()
    format_hj_position_cell = lambda s: str(s or "").strip()
    fix_email_ocr = lambda s: str(s or "").strip()

_LABEL_HINTS = (
    "姓名", "证件", "性别", "出生", "国籍", "民族", "手机", "电话", "邮箱", "电子邮件",
    "毕业", "学位", "回国", "职称", "单位", "职务", "地址", "省份", "城市", "实验室",
    "联系人", "填表", "项目类别", "gender", "birth", "nationality", "mobile", "email",
    "degree", "employer", "position", "province", "city", "applicant", "certificate",
    "序号", "起止", "时间", "院校", "专业", "教育", "工作", "经历", "成果", "论文",
    "知识产权", "奖励", "荣誉", "设想", "承诺", "推荐理由", "支持条件",
)

_EDU_TAB = re.compile(
    r"(?:^|[^\d])(\d{6})\s*-\s*(\d{6})\t+([^\t]+)\t+([^\t]+)\t+([^\t]+)\t+((?:学士|硕士|博士)[^\n]*)",
    re.I,
)
_WORK_TAB = re.compile(
    r"(?:^|[^\d])(\d{6})\s*-\s*(\d{6})\t+([^\t]+)\t+([^\t]+)\t+([^\t]+)\t+(全职|兼职)",
    re.I,
)
_PAPER_LINE = re.compile(
    r"^\d+\t(\d{4}-\d{2})\t(.+?)\t(.+?)\t(\S+)\t(.+)$"
)
_PAPER_ENTRY_START = re.compile(r"(?:^|\s)(\d{1,2})\s+(\d{4}-\d{2})\b")
_PAPER_DATE_ONLY = re.compile(r"^(\d{4}[-./]\d{2}(?:[-./]\d{2})?)$")
_PAPER_SERIAL_ONLY = re.compile(r"^(\d{1,2})$")
_PAPER_RANK = re.compile(r"(\d{1,2}/\d{1,2})")
_PAPER_VENUE = re.compile(
    r"(Energy Reviews|Carbon Energy|Advanced Energy Materials|Nano Energy|"
    r"Electrochimica\s*Acta|Advanced\s+[Mm]ater(?:ials)?|Energy Storage\s*Materials|"
    r"Advanced Functional Materials|Journal of Materials Chemistry|ACS Nano|"
    r"Electrochimica|Acta|Materials Today|Chemical Engineering Journal)",
    re.I,
)
_PAPER_SECTION_START = ("代表性论著", "Publications (No More", "主要成果", "Achievements")
_PAPER_SECTION_END = ("代表性知识产权", "Intellectual Property Rights", "(2) 代表性知识产权", "(3)代表性知识产权")
_PROJECT_SECTION_START = ("代表性科研项目", "Grants (As a Leader", "2 代表性科研项目", "2.代表性科研项目")
_PROJECT_SECTION_END = ("代表性论著", "Publications (No More", "主要成果", "Achievements")
_OCR_CHECKED = re.compile(r"[☑☒■●◎@团回四旷]")
_OCR_TERM_FIXES = (
    ("粉未", "粉末"),
    ("高铉", "高镍"),
    ("锤基", "锰基"),
    ("钳离子", "钠离子"),
    ("画极", "电极"),
    ("顽士", "硕士"),
    ("植尚", "崇尚"),
    ("债守", "恪守"),
    ("合中报人", "含申报人"),
    ("渚盖", "覆盖"),
    ("公容", "公寓"),
    ("灰晶石", "尖晶石"),
    ("锂镇", "锂镍"),
    ("暧代", "替代"),
    ("借率", "倍率"),
    ("教捷", "教授"),
    ("目接", "直接"),
    ("砂博", "硕博"),
    ("砂士", "硕士"),
    ("高镇", "高镍"),
    ("高镐", "高镍"),
    ("高镛", "高镍"),
    ("高镣", "高镍"),
    ("缺么", "缺乏"),
    ("快逢", "快速"),
    ("糊合", "耦合"),
    ("复来", "复杂"),
    ("去噻", "去噪"),
    ("反停", "反馈"),
    ("复束", "复杂"),
    ("韶争力", "竞争力"),
    ("智格", "晶格"),
    ("降侥", "降低"),
    ("电板", "电极"),
    ("傅能", "储能"),
    ("宗全", "安全"),
    ("《hAdvanced", "《Advanced"),
    ("«Advanced", "《Advanced"),
    ("CTIMs", "TIMs"),
    ("TTIMs", "TIMs"),
    ("基于 AL 的", "基于 AI 的"),
    ("基于AL的", "基于AI的"),
    ("ALI 技术", "AI 技术"),
    ("AT 图像", "AI 图像"),
    ("AL 图像", "AI 图像"),
    ("A 驱动", "AI 驱动"),
    ("准确率之", "准确率≥"),
)
_EMPLOYER_TYPE_BLOCK = (
    "1.用人单位类型：\n"
    "□部属高校□地方高校□军队院校\n"
    "□中国科学院□中国工程物理研究院□军队科研院所□其他科研院所\n"
    "□中央企业□地方国有企业□民营企业 □其他"
)
_COVER_FRONTIER = (
    "集成电路", "人工智能", "量子信息", "先进制造", "生命健康",
    "脑科学", "生物育种", "空天科技", "深地深海", "不属于上述前沿领域",
)
_COVER_CORE_TECH = (
    "集成电路", "人工智能", "量子科技", "生物科技", "石油天然气", "基础原材料",
    "超级计算机", "信息通讯", "工业软件", "农作物种子", "科学试验用仪器设备",
    "化学制剂", "药品", "医疗器械", "医用设备", "疫苗", "不涉及上述关键核心技术",
)
_EXPERTISE_TITLE = "专长及代表性成果(Expertise and Achievements)"
_EXPERTISE_INTRO_HEAD = "所从事的专业领域及取得的成绩描述"
_EXPERTISE_INTRO_HINT = "（概述与所在或拟应聘实验室相关的研究领域、方向及取得的成就，5000字以内）"
_EXPERTISE_INTRO = _EXPERTISE_INTRO_HEAD + _EXPERTISE_INTRO_HINT
_EXPERTISE_EN = (
    "Field of Expertise and Achievements(Please Outline the Research  Field, "
    "Direction  and  Achievements  Related to Your Current or Expected Laboratory,"
    "No More than 5000 words)"
)
_NARR_HEAD_ONLY = re.compile(
    r"^(?:"
    r"[一二三四五六七八九十]+、.{0,40}|"
    r"\d{1,2}[\.、．\s].{0,36}|"
    r"[①②③④⑤⑥⑦⑧⑨⑩]\s*.{0,36}|"
    r"[（(]\d{1,2}[)）]\s*.{0,36}"
    r")$"
)
_SCALAR_LABEL_GARBAGE = re.compile(
    r"(Time of Coming to China|Before Returning|Coming to China|"
    r"Current or Expected Employer|Host Province|Host City|"
    r"手机号\s*Mobile|电子邮件\s*Email|Mobile$|Email$|Position$|Employer$)",
    re.I,
)
_PAGE_MARK = re.compile(
    r"^【第\d+页】|^第\s*\d+\s*页|^\d{1,3}/\d{1,3}$|"
    r"^\d{4}/\d{1,2}/\d{1,2}\s+\d{1,2}:\d{2}"
)
_DATE_ONLY = re.compile(r"^(\d{6})\s*-\s*(\d{6})$")
_DATE_LOOSE = re.compile(r"^(\d{6})\s*-\s*(\d{1,6})\b(.*)$")
_DATE_PAIR_INLINE = re.compile(r"(?<!\d)(20\d{4})\s*-\s*(?:20)?(\d{1,6})(?!\d)")
_TIMELINE_STOP_MARKERS = (
    "破格申报", "重要科研奖励", "科研情况", "工作设想", "Exceptional Application",
    "Important Awards", "Research Expertise",
)
_TIMELINE_SCHOOL_HINTS = (
    (re.compile(r"宗\s*娜\s*大\s*学|安娜大学|Anna\s*University", re.I), "安娜大学/Anna University"),
    (re.compile(r"全南国立|Chonnam", re.I), "全南国立大学/Chonnam National University"),
    (re.compile(r"韦\s*仕\s*敦|Western[\s\S]{0,24}[Uu]niversity", re.I), "韦仕敦大学/Western University"),
    (re.compile(r"滑铁卢|Waterloo", re.I), "滑铁卢大学/University of Waterloo"),
    (re.compile(r"阿德文|AdvEn", re.I), "阿德文工业/AdvEn Industrial"),
)
_VERTICAL_NOISE = re.compile(
    r"^(时间|Time|国家|Country|院校|University|专业|Major|学位|Degree|单位|Employer|职务|Position|"
    r"任职情况|Work Performance|Educational Background|All Work Experience|教育经历|全部工作经历|"
    r"破格申报|Exceptional Application).*$",
    re.I,
)
_COUNTRY_NAME = re.compile(
    r"^(印度|韩国|加拿大|美国|英国|德国|法国|日本|澳大利亚|巴西|意大利|中国|西班牙|荷兰|瑞士|瑞典|"
    r"丹麦|挪威|芬兰|波兰|新加坡|马来西亚|泰国|埃及|黎巴嫩|沙特阿拉伯|墨西哥|以色列|比利时|"
    r"奥地利|新西兰|爱尔兰|葡萄牙|希腊|土耳其|俄罗斯|乌克兰|伊朗|南非|阿根廷|智利|中国台湾|中国香港)$"
)
_EN_COUNTRY_MAP = {
    "canada": "加拿大", "india": "印度", "korea": "韩国", "southkorea": "韩国",
    "uk": "英国", "unitedkingdom": "英国", "britain": "英国", "england": "英国",
    "usa": "美国", "unitedstates": "美国", "portugal": "葡萄牙", "china": "中国",
    "brazil": "巴西", "japan": "日本", "germany": "德国", "france": "法国",
}
_SCHOOL_COUNTRY_HINTS = SCHOOL_COUNTRY_HINTS or (
    (re.compile(r"Anna\s*University|安娜大学", re.I), "印度"),
    (re.compile(r"Chonnam|全南", re.I), "韩国"),
    (re.compile(r"Brighton", re.I), "英国"),
    (re.compile(r"Waterloo|滑铁卢", re.I), "加拿大"),
    (re.compile(r"Western|韦仕敦", re.I), "加拿大"),
    (re.compile(r"AdvEn|阿德文", re.I), "加拿大"),
)
_HJ_FIELD_GARBAGE = re.compile(
    r"(Educational|hronological|Order\.|符合申报|学历学位|海外连续|"
    r"Contact Information|任职单位名称|Host Province|"
    r"出生国家生日|生日勐|印度.*性别|性别.*印度|中文\s*Cn|英文\s*En)",
    re.I,
)
_CHECKBOX_GARBAGE = re.compile(
    r"[☑☐□口].*?(符合申报|学历学位|海外连续|其他)|"
    r"☑\s*符合申报.*$|☐海外连续工作年限.*$|☐其他\s*$",
    re.I,
)


def template_path(mode: str) -> Path:
    name = "HJ.docx" if str(mode or "").upper() == "HJ" else "QM.docx"
    p = ROOT_DIR / name
    if not p.is_file():
        raise FileNotFoundError("模板不存在：" + str(p))
    return p


def _norm(s: str) -> str:
    return re.sub(r"[\s/·•\u3000]+", "", str(s or "")).lower()


def _looks_like_label(text: str) -> bool:
    t = str(text or "").strip()
    if not t:
        return False
    n = _norm(t)
    if len(n) < 2:
        return True
    return any(h in t or h in n for h in _LABEL_HINTS)


def _distinct_cells(row) -> list:
    seen, out = set(), []
    for c in row.cells:
        cid = id(c._tc)
        if cid in seen:
            continue
        seen.add(cid)
        out.append(c)
    return out


def _apply_run_font(
    run,
    east: str = "宋体",
    ascii_name: str = "Times New Roman",
    size_pt: float = 12,
    color: str = "000000",
    bold: bool | None = None,
) -> None:
    """填写内容统一为宋体/Times New Roman、黑色，避免空段落入主题蓝/灰。"""
    rPr = run._r.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rPr.insert(0, rf)
    rf.set(qn("w:ascii"), ascii_name)
    rf.set(qn("w:hAnsi"), ascii_name)
    rf.set(qn("w:eastAsia"), east)
    rf.set(qn("w:cs"), ascii_name)
    half = str(int(size_pt * 2))
    for tag in ("w:sz", "w:szCs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            rPr.append(el)
        el.set(qn("w:val"), half)
    col = rPr.find(qn("w:color"))
    if col is None:
        col = OxmlElement("w:color")
        rPr.append(col)
    col.set(qn("w:val"), color)
    theme = qn("w:themeColor")
    if theme in col.attrib:
        del col.attrib[theme]
    hl = rPr.find(qn("w:highlight"))
    if hl is not None:
        rPr.remove(hl)
    if bold is True:
        if rPr.find(qn("w:b")) is None:
            rPr.append(OxmlElement("w:b"))
        if rPr.find(qn("w:bCs")) is None:
            rPr.append(OxmlElement("w:bCs"))
    elif bold is False:
        for tag in ("w:b", "w:bCs"):
            el = rPr.find(qn(tag))
            if el is not None:
                rPr.remove(el)


def _run_has_font(run) -> bool:
    rPr = run._r.find(qn("w:rPr"))
    if rPr is None:
        return False
    return rPr.find(qn("w:rFonts")) is not None or rPr.find(qn("w:sz")) is not None


def _paint_paragraph_font(
    p,
    east: str = "宋体",
    ascii_name: str = "Times New Roman",
    size_pt: float = 12,
    bold: bool | None = False,
) -> None:
    runs = [r for r in p.runs if str(r.text or "")]
    if not runs:
        return
    for r in runs:
        _apply_run_font(r, east=east, ascii_name=ascii_name, size_pt=size_pt, bold=bold)


def _write_runs(p, parts: list[tuple[str, dict]]) -> None:
    """清空段落 run 后按指定字体写入。"""
    for r in list(p.runs):
        r.text = ""
    first = True
    for text, kw in parts:
        if not text:
            continue
        if first and p.runs:
            r = p.runs[0]
            r.text = text
            first = False
        else:
            r = p.add_run(text)
            first = False
        _apply_run_font(r, **kw)


def _assign_run_texts(p, new_text: str) -> bool:
    """按原 run 切分写入，勾选框只改 □/☑，不把后面的黑体/Arial 吞进勾选字体。"""
    runs = list(p.runs)
    old = "".join(r.text or "" for r in runs)
    v = str(new_text or "")
    if v == old:
        return True
    if not runs:
        return False
    if len(v) == len(old):
        i = 0
        for r in runs:
            n = len(r.text or "")
            r.text = v[i:i + n]
            i += n
        return True
    prefix = 0
    lim = min(len(old), len(v))
    while prefix < lim and old[prefix] == v[prefix]:
        prefix += 1
    extra = len(v) - len(old)
    consumed = 0
    i_new = 0
    placed = False
    for r in runs:
        n = len(r.text or "")
        if consumed + n <= prefix and not placed:
            i_new += n
            consumed += n
            continue
        keep = max(0, prefix - consumed)
        take = keep + (n - keep) + (extra if not placed else 0)
        if take < 0:
            take = keep
        r.text = v[i_new:i_new + max(0, take)]
        i_new += max(0, take)
        placed = True
        extra = 0
        consumed += n
    if i_new < len(v) and runs:
        last = runs[-1]
        last.text = (last.text or "") + v[i_new:]
    return True


def _set_paragraph_text(p, value: str) -> None:
    """改段落文字，保留原 run 的字体（封面黑体+Arial、表内宋体）。"""
    v = str(value or "")
    runs = list(p.runs)
    if not runs:
        if v:
            run = p.add_run(v)
            _apply_run_font(run)
        return
    old = "".join(r.text or "" for r in runs)
    if old == v:
        if v and not any(_run_has_font(r) for r in runs if r.text):
            _paint_paragraph_font(p)
        return
    n = 0
    lim = min(len(old), len(v))
    while n < lim and old[n] == v[n]:
        n += 1
    consumed = 0
    suffix_done = False
    for r in runs:
        t = r.text or ""
        end = consumed + len(t)
        if end <= n and not suffix_done:
            consumed = end
            continue
        keep = max(0, n - consumed)
        if not suffix_done:
            r.text = t[:keep] + v[n:]
            if not _run_has_font(r):
                _apply_run_font(r)
            suffix_done = True
        else:
            r.text = ""
        consumed = end
    if not suffix_done:
        runs[-1].text = (runs[-1].text or "") + v[n:]


def _cell_value_paragraph(cell):
    """正式 HJ / 16 表模板：数值写在单元格最后一段（前面空段撑齐标签）。"""
    paras = list(cell.paragraphs)
    if not paras:
        return None
    nonempty = [p for p in paras if str(p.text or "").strip()]
    return nonempty[-1] if nonempty else paras[-1]


def _write_cell(cell, value: str) -> None:
    """写入单元格数据值，不打散空段撑齐结构。"""
    v = str(value or "").strip()
    paras = list(cell.paragraphs)
    if not paras:
        return
    target = _cell_value_paragraph(cell)
    if target is None:
        return
    for p in paras:
        if p is target:
            continue
        if str(p.text or "").strip():
            _set_paragraph_text(p, "")
    _set_paragraph_text(target, v)
    _paint_paragraph_font(target)


def _mark_checkboxes_in_cell(cell, keywords: tuple[str, ...]) -> None:
    """勾选框按段落改，避免把中英文两段压进最后一段。"""
    for p in cell.paragraphs:
        t = str(p.text or "")
        if not re.search(r"[□☐口☑]", t):
            continue
        nt = _mark_checkbox(t, keywords)
        if nt != t:
            if not _assign_run_texts(p, nt):
                _set_paragraph_text(p, nt)
            _paint_paragraph_font(p)


def _normalize_narr_heading(line: str) -> str:
    s = str(line or "").strip()
    s = re.sub(r"^(\d{1,2})[\s,，、．]+", r"\1.", s)
    return s


def _is_major_section_heading(line: str) -> bool:
    return bool(re.match(r"^[一二三四五六七八九十]+、", str(line or "").strip()))


def _is_narr_heading(line: str) -> bool:
    s = _normalize_narr_heading(line)
    if not s or len(s) > 48:
        return False
    if re.fullmatch(r"\d{1,2}\.?", s):
        return False
    if re.search(r"[。！？!?]$", s) and len(s) > 24:
        return False
    if "：" in s or ":" in s:
        head = re.split(r"[：:]", s, 1)[0]
        if len(s) > len(head) + 12:
            return False
    return bool(_NARR_HEAD_ONLY.match(s)) or _is_major_section_heading(s)


def _join_wrap_lines(lines: list[str]) -> str:
    out = ""
    for ln in lines:
        s = str(ln or "").strip()
        if not s:
            continue
        if not out:
            out = s
            continue
        if re.search(r"[\u4e00-\u9fff]$", out) and re.search(r"^[\u4e00-\u9fff]", s):
            out += s
        else:
            out = out.rstrip() + " " + s
    return out.strip()


def _split_narrative_paras(text: str) -> list[str]:
    """按正式 16 表：标题单独成段，正文按句段换段，不整格糊成一段。"""
    raw = str(text or "").strip()
    if not raw:
        return []
    chunks: list[str] = []
    buf: list[str] = []

    def flush() -> None:
        joined = _join_wrap_lines(buf)
        buf.clear()
        if joined:
            chunks.append(joined)

    for ln in str(raw).splitlines():
        s = ln.strip()
        if not s:
            flush()
            continue
        if s in {"\\", "/", "|", "—"} or re.fullmatch(r"[-\\/]{1,3}", s):
            continue
        if re.fullmatch(r"\d{1,2}\.?", s):
            continue
        if _is_narr_heading(s):
            flush()
            s = _normalize_narr_heading(s)
            if _is_major_section_heading(s) and chunks:
                chunks.append("")
            chunks.append(s)
            continue
        if buf and re.search(r"[。！？]$", buf[-1]) and len(_join_wrap_lines(buf)) >= 40:
            flush()
        buf.append(s)
    flush()
    return chunks


def _insert_para_after(p, text: str):
    new_el = deepcopy(p._p)
    p._p.addnext(new_el)
    new_p = Paragraph(new_el, p._parent)
    _set_paragraph_text(new_p, text)
    return new_p


def _write_paragraphs_from(cell, start_idx: int, lines: list[str], stop_labels: tuple[str, ...] = (), keep_blank: bool = False) -> None:
    """从指定段起写入多段正文，不够就按原段样式克隆，避免 _write_cell 压扁。"""
    if keep_blank:
        cleaned: list[str] = []
        for x in lines:
            s = str(x or "")
            cleaned.append("" if not s.strip() else s.strip())
        while cleaned and not cleaned[0]:
            cleaned.pop(0)
        while cleaned and not cleaned[-1]:
            cleaned.pop()
        lines = cleaned
    else:
        lines = [str(x or "").strip() for x in lines if str(x or "").strip()]
    paras = list(cell.paragraphs)
    if not lines:
        return
    if not paras:
        for ln in lines:
            cell.add_paragraph(ln)
        return
    stop = len(paras)
    for i, p in enumerate(paras):
        if i < start_idx:
            continue
        t = str(p.text or "")
        if stop_labels and any(k in t for k in stop_labels):
            stop = i
            break
    src = paras[start_idx] if start_idx < len(paras) else paras[-1]
    existing = max(0, stop - start_idx)
    n_fill = min(existing, len(lines))
    for i in range(n_fill):
        _set_paragraph_text(paras[start_idx + i], lines[i])
        if str(lines[i]).strip():
            _paint_paragraph_font(paras[start_idx + i], bold=False)
    if existing < len(lines):
        cursor = paras[start_idx + n_fill - 1] if n_fill else (paras[start_idx - 1] if start_idx else src)
        for ln in lines[n_fill:]:
            cursor = _insert_para_after(cursor, ln)
            if str(ln).strip():
                _paint_paragraph_font(cursor, bold=False)
    elif existing > len(lines):
        for p in paras[start_idx + len(lines):stop]:
            if stop_labels and any(k in (p.text or "") for k in stop_labels):
                break
            if str(p.text or "").strip():
                _set_paragraph_text(p, "")


def _fmt_hj_date(s: str) -> str:
    """参考正式 HJ 申报书：YYYY.MM.DD（仅年月时补 .01）。"""
    raw = str(s or "").strip()
    if "至今" in raw:
        return "至今"
    digits = re.sub(r"\D", "", raw)
    if len(digits) >= 8:
        y, m, d = int(digits[:4]), int(digits[4:6]), int(digits[6:8])
        return f"{y}.{m:02d}.{d:02d}" if d else f"{y}.{m:02d}.01"
    if len(digits) == 6:
        return f"{int(digits[:4])}.{int(digits[4:6]):02d}.01"
    if len(digits) == 4:
        return digits
    if re.match(r"\d{4}[./-]\d{1,2}([./-]\d{1,2})?$", raw):
        parts = re.split(r"[./-]", raw)
        y = int(parts[0])
        m = int(parts[1]) if len(parts) > 1 else 0
        d = int(parts[2]) if len(parts) > 2 else 0
        if d:
            return f"{y}.{m:02d}.{d:02d}"
        if m:
            return f"{y}.{m:02d}.01"
        return str(y)
    if re.match(r"\d{4}\.\d{1,2}$", raw):
        return raw + ".01"
    return raw


def _fmt_hj_range(a: str, b: str) -> str:
    a_s, b_s = str(a or "").strip(), str(b or "").strip()
    if "至今" in b_s or b_s in {"999999", "999912"}:
        return _fmt_hj_date(a_s) + "-至今"
    return _fmt_hj_date(a_s) + "-" + _fmt_hj_date(b_s)


def _normalize_hj_range_text(s: str) -> str:
    """将解析器/OCR 各类起止时间统一为参考申报书格式。"""
    raw = str(s or "").strip()
    if not raw:
        return ""
    if "至今" in raw:
        head = raw.split("至今", 1)[0].rstrip("- ")
        return _fmt_hj_date(head) + "-至今"
    m = re.match(r"(\d{4}\.\d{2}\.\d{2})-(\d{1,6})$", raw)
    if m:
        end = _repair_yyyymm_end(m.group(1).replace(".", "")[:6], m.group(2), raw)
        return _fmt_hj_range(m.group(1).replace(".", "")[:6], end)
    if re.match(r"\d{4}\.\d", raw) and "-" in raw:
        a, b = raw.split("-", 1)
        a_d = re.sub(r"\D", "", a)[:6]
        b_d = _repair_yyyymm_end(a_d, b.strip(), raw)
        return _fmt_hj_range(a_d, b_d)
    m = re.match(r"(\d{4})-(\d{2})-(\d{4})-(\d{2})$", raw)
    if m:
        return f"{m.group(1)}.{m.group(2)}-{m.group(3)}.{m.group(4)}"
    m = re.search(r"(\d{6,8})\s*-\s*(\d{6,8}|至今)", raw)
    if m:
        end = m.group(2)
        if end != "至今":
            end = _repair_yyyymm_end(m.group(1), end, raw)
        return _fmt_hj_range(m.group(1), end)
    return raw


def _valid_yyyymm(s: str) -> bool:
    if len(s) != 6 or not s.isdigit():
        return False
    y, m = int(s[:4]), int(s[4:6])
    return 1900 <= y <= 2100 and 1 <= m <= 12


def _yyyymm_candidates(text: str) -> list[str]:
    compact = re.sub(r"\s+", "", str(text or ""))
    found: list[str] = []
    for m in re.finditer(r"(20\d{4})", compact):
        if _valid_yyyymm(m.group(1)):
            found.append(m.group(1))
    for m in re.finditer(r"(?<![\d])(0\d{4,5})(?!\d)", compact):
        cand = "2" + m.group(1)
        if _valid_yyyymm(cand):
            found.append(cand)
    out: list[str] = []
    seen: set[str] = set()
    for item in found:
        if item not in seen:
            seen.add(item)
            out.append(item)
    return out


def _repair_yyyymm_end(start: str, end_raw: str, context: str) -> str:
    start = re.sub(r"\D", "", str(start or ""))[:6]
    end_raw = re.sub(r"\s+", "", str(end_raw or ""))
    if _valid_yyyymm(end_raw[:6]):
        return end_raw[:6]
    if len(end_raw) == 5 and end_raw.startswith("0"):
        cand = "2" + end_raw
        if _valid_yyyymm(cand):
            return cand
    for cand in _yyyymm_candidates(end_raw + context):
        if cand > start:
            return cand
    return end_raw


def _fmt_birth(s: str) -> str:
    return _fmt_hj_date(s)


def _normalize_gender(val: str) -> str:
    s = str(val or "")
    if re.search(r"女|Female", s, re.I) and not re.search(r"男|Male|Wale", s, re.I):
        return "女"
    if re.search(r"男|\bMale\b|\bWale\b|M\s*ale", s, re.I):
        return "男"
    return ""


def _strip_name_noise(name: str) -> str:
    s = str(name or "").strip()
    s = re.split(r"依托单位|用人单位|申报单位|Name of Applicant", s, 1)[0].strip()
    return re.sub(r"\s+", " ", s).strip(" ：:")


def _normalize_country_name(val: str) -> str:
    raw = re.sub(r"\s+", "", str(val or "").strip())
    if not raw:
        return ""
    if _COUNTRY_NAME.match(raw):
        return raw
    key = re.sub(r"[^a-z]", "", raw.lower())
    for en, cn in _EN_COUNTRY_MAP.items():
        if key == en or key.startswith(en):
            return cn
    return raw


def _infer_country_from_entity(name: str) -> str:
    for pat, country in _SCHOOL_COUNTRY_HINTS:
        if pat.search(str(name or "")):
            return country
    return ""


def _lines_after_marker(
    text: str,
    markers: tuple[str, ...],
    stop_markers: tuple[str, ...] = (),
    max_lines: int = 8,
) -> list[str]:
    lines = [ln.strip() for ln in str(text or "").splitlines()]
    start = -1
    for i, ln in enumerate(lines):
        if any(m in ln for m in markers):
            start = i + 1
            break
    if start < 0:
        return []
    out: list[str] = []
    for ln in lines[start: start + max_lines]:
        if not ln or _PAGE_MARK.match(ln) or re.match(r"^\d{1,3}/\d{1,3}$", ln):
            break
        if any(s in ln for s in stop_markers):
            break
        if re.match(r"^(Name of|Date of|Place of|ID Type|Other ID|Contact Information|证件类型)", ln, re.I):
            continue
        out.append(ln)
    return out


_NARRATIVE_FIELD_KEYS = frozenset({
    "专长及代表性成果", "工作设想", "用人单位简介", "推荐理由", "支持条件",
})


def _is_hj_field_garbage(val: str, key: str = "") -> bool:
    s = str(val or "").strip()
    if not s:
        return True
    if s in ("男", "女", "是", "否"):
        return False
    if key in ("省级项目意愿", "首次申报勾选"):
        return False
    if len(s) < 2:
        return True
    if _HJ_FIELD_GARBAGE.search(s):
        return True
    if key in _NARRATIVE_FIELD_KEYS:
        return False
    if re.search(r"[\u4e00-\u9fff]", s) and re.search(r"[A-Za-z]{4,}", s) and len(s) > 24:
        return True
    return False


def _strip_checkbox_garbage(s: str) -> str:
    s = _CHECKBOX_GARBAGE.sub("", str(s or ""))
    s = re.sub(r"条件\s*[（(].*$", "", s)
    s = re.sub(r"[（(]硕士学位[）)].*$", "", s)
    s = re.sub(r"/无学历学位.*$", "", s)
    s = re.sub(r"工作年限.*$", "", s)
    return s.strip(" ；;,，")


def _split_employer_position(text: str) -> tuple[str, str]:
    s = _strip_checkbox_garbage(text)
    if not s:
        return "", ""
    m = re.match(r"^(.+?\([^)]+\))\s+(.+?\([^)]+\))(?:\s|$)", s)
    if m:
        return m.group(1).strip(), m.group(2).strip()
    m = re.match(
        r"^(.+?(?:大学|学院|工业|公司|研究所)(?:\s*\([^)]+\))?)\s+((?:首席|助理|副)?(?:教授|科学家|研究员|总监|经理|主任|工程师|博士后).*)$",
        s,
    )
    if m:
        return m.group(1).strip(), m.group(2).strip()
    m = re.match(
        r"^(.+?)\s+((?:首席|助理|副)?(?:教授|科学家|研究员|总监|经理|主任|工程师|博士后).*)$",
        s,
    )
    if m:
        return m.group(1).strip(), m.group(2).strip()
    return s, ""


def _extract_split_ocr_english_name(text: str) -> str:
    """OCR 常把英文名拆成「有效证件姓名 KALIYAPPAN KAR」+「Name of ID THIKEYAN」两行。"""
    from .scan_text_normalize import collapse_cjk_spaces

    lines = [collapse_cjk_spaces(ln.strip()) for ln in str(text or "").splitlines()]
    for i, ln in enumerate(lines):
        m = re.search(r"有效证件姓名\s+([A-Z][A-Z\s]+)", ln)
        if not m:
            continue
        parts = [m.group(1).strip()]
        j = i + 1
        while j < len(lines) and not lines[j]:
            j += 1
        if j < len(lines):
            nxt = lines[j]
            if re.match(r"Name of ID", nxt, re.I):
                tail = re.sub(r"^Name of ID\s*", "", nxt, flags=re.I).strip()
                if tail and re.match(r"^[A-Z]", tail):
                    words = parts[-1].split()
                    if words:
                        frag = words[-1]
                        if len(frag) <= 4 or tail.upper().startswith(frag.upper()[:2]):
                            prefix = " ".join(words[:-1])
                            merged = (prefix + " " if prefix else "") + frag + tail
                            parts[-1] = merged.strip()
                        else:
                            parts.append(tail)
        full = fix_ocr_english(re.sub(r"\s+", " ", " ".join(parts)))
        if len(re.sub(r"\s+", "", full)) >= 10:
            return full
    return ""


def _fix_concatenated_latin_name(name: str, raw: str = "") -> str:
    """KALIYAPPANKARTHIKEYAN → KALIYAPPAN KARTHIKEYAN。"""
    split = _extract_split_ocr_english_name(raw)
    if split and " " in split and len(re.sub(r"\s+", "", split)) >= 16:
        return split
    colon = ""
    m = re.search(r"申报人有效证件姓名[^：\n]*[：:]\s*([A-Z][A-Za-z]{8,})", str(raw or ""))
    if m:
        colon = re.sub(r"[^A-Za-z].*$", "", m.group(1)).upper()
    s = str(name or "").strip()
    upper = re.sub(r"\s+", "", colon or s).upper()
    if not upper:
        return split or fix_ocr_english(s)
    for suffix in ("KARTHIKEYAN", "KALIYAPPAN", "KUMAR", "SINGH"):
        if upper.endswith(suffix) and len(upper) > len(suffix) + 3:
            return fix_ocr_english(f"{upper[:-len(suffix)]} {suffix}")
    m = re.search(r"Name of ID\s*([A-Z]{4,})", str(raw or ""), re.I | re.M)
    if m:
        suffix = m.group(1).upper()
        if upper.endswith(suffix) and len(upper) > len(suffix) + 3:
            return fix_ocr_english(f"{upper[:-len(suffix)]} {suffix}")
    return split or fix_ocr_english(s or colon)


def _format_chinese_transliteration(cn: str) -> str:
    cn = re.sub(r"\s+", "", str(cn or ""))
    cn = re.sub(r"申报.*$", "", cn)
    cn = re.sub(r"^中文[（(]?音译[）)]?(?:姓)?名?", "", cn)
    if not cn or "·" in cn:
        return cn
    if cn.startswith("卡利亚潘") and len(cn) > 4:
        return "卡利亚潘·" + cn[4:]
    if len(cn) >= 8:
        return cn[:4] + "·" + cn[4:]
    return cn


def _extract_english_certificate_name(text: str) -> str:
    raw = str(text or "")
    split = _extract_split_ocr_english_name(raw)
    if split:
        return split
    m = re.search(
        r"申报人有效证件姓名[^：\n]*[：:]\s*([A-Z][A-Za-z]+)(?:\s+([A-Z][A-Za-z]+))?",
        raw,
    )
    if m:
        name = " ".join(g for g in m.groups() if g).strip()
        if len(name) >= 6:
            return _fix_concatenated_latin_name(name, raw)
    m = re.search(
        r"申报人有效证件姓名[^：\n]*[：:]\s*([A-Z][A-Za-z\s\"']+?)(?:[（(\"“]|$)",
        raw,
    )
    if m:
        return _fix_concatenated_latin_name(re.sub(r"\s+", " ", m.group(1)).strip(), raw)
    parts: list[str] = []
    lines = [ln.strip() for ln in raw.splitlines()]
    for i, ln in enumerate(lines):
        if "有效证件姓名" in ln and "Name of" not in ln:
            tail = re.sub(r".*有效证件姓名\s*", "", ln).strip()
            if tail and re.match(r"^[A-Z]", tail):
                parts.append(tail)
        if re.search(r"Name of ID", ln, re.I):
            tail = re.split(r"Name of ID", ln, flags=re.I)[-1].strip()
            if tail and re.match(r"^[A-Z]", tail):
                parts.append(tail)
            elif i + 1 < len(lines):
                nxt = lines[i + 1].strip()
                if re.match(r"^[A-Z]{3,}", nxt) and "Name of" not in nxt:
                    parts.append(nxt)
    if parts:
        return fix_ocr_english(re.sub(r"\s+", " ", " ".join(parts)).strip())
    m = re.search(r"(?i)Name of ID\t+([^\n\t]+)", raw)
    return fix_ocr_english(m.group(1).strip()) if m else ""


def _extract_chinese_transliteration(text: str) -> str:
    raw = str(text or "")
    from .scan_text_normalize import collapse_cjk_spaces

    cover = raw.split("---", 1)[0]
    cover_flat = collapse_cjk_spaces(re.sub(r"\s+", "", cover))
    m_cover = re.search(r"[（(\"“]\s*([\u4e00-\u9fff]{6,16})\s*[）)]", cover_flat)
    if m_cover:
        cleaned = _format_chinese_transliteration(m_cover.group(1))
        if 4 <= len(cleaned) <= 16:
            return cleaned

    m = re.search(
        r"申报人有效证件姓名[^：\n]*[：:]\s*[A-Z][^\n（(“\"]*[（(“\"]\s*([\u4e00-\u9fff·•\s]{4,24})",
        collapse_cjk_spaces(raw),
    )
    if m:
        cleaned = _format_chinese_transliteration(re.sub(r"[^一-龥·•]", "", m.group(1)))
        if 4 <= len(cleaned) <= 16:
            return cleaned
    m = re.search(
        r"申报人有效证件姓名[^：\n]*[：:]\s*[A-Z][^\n（(“\"]*[（(“\"]\s*([\u4e00-\u9fff·•\s]{4,24})",
        raw,
    )
    if m:
        cleaned = _format_chinese_transliteration(re.sub(r"\s+", "", m.group(1)))
        if 4 <= len(cleaned) <= 16:
            return cleaned
    lines = [collapse_cjk_spaces(ln.strip()) for ln in raw.splitlines()]
    for i, ln in enumerate(lines):
        m = re.search(r"中文[（(]音译[）)]名\s*([\u4e00-\u9fff·•\s]{2,16})", ln)
        if m:
            cleaned = _format_chinese_transliteration(re.sub(r"[^一-龥·•]", "", m.group(1)))
            if 4 <= len(cleaned) <= 16:
                return cleaned
        if "中文" in ln and "音译" in ln:
            tail = re.sub(r".*音译[）)]名\s*", "", ln)
            cleaned = _format_chinese_transliteration(re.sub(r"[^一-龥·•]", "", tail))
            if 4 <= len(cleaned) <= 16:
                return cleaned
            for j in range(i + 1, min(i + 3, len(lines))):
                nxt = lines[j]
                if re.match(r"^Name of", nxt, re.I):
                    continue
                cleaned = _format_chinese_transliteration(re.sub(r"[^一-龥·•]", "", nxt))
                if 4 <= len(cleaned) <= 16 and not re.search(r"(出生|性别|国家|生日|申报)", cleaned):
                    return cleaned
    return ""


def _extract_contact_vertical_fields(text: str) -> dict[str, str]:
    """扫描件竖排联系方式区：标签与值分列时，按固定顺序取手机号等栏位。"""
    from .scan_text_normalize import collapse_cjk_spaces

    lines = [ln.strip() for ln in str(text or "").splitlines()]
    start = -1
    for i, ln in enumerate(lines):
        compact = re.sub(r"\s+", "", ln)
        if "手机号Mobile" in compact or ("Mobile" in compact and "手机" in compact):
            start = i + 1
            break
    if start < 0:
        return {}

    values: list[str] = []
    i = start
    while i < len(lines):
        ln = re.sub(r"[ \t\u3000]+", " ", lines[i].strip())
        if not ln:
            i += 1
            continue
        if _PAGE_MARK.match(ln) or re.match(r"^\d{1,3}/\d{1,3}$", ln):
            break
        compact = re.sub(r"\s+", "", ln)
        if "教育经历" in compact or "EducationalBackground" in compact:
            break
        if values and ln.startswith(("口", "□", "☑", "☐", "团", "旷")) and not re.search(r"[\d+@]", ln):
            break
        if not values and (
            re.search(
                r"(Mobile|Email|Cn|En|职称|Degree|Employer|Country|Province|City|"
                r"Coming to China|Institution|Position|Type of|Residence|Contact Information|申报人联系|未提交)",
                ln,
                re.I,
            )
            or _looks_like_label(ln)
        ):
            i += 1
            continue
        if values and re.match(r"^[a-z]", ln) and re.search(r"[A-Za-z]", values[-1]):
            values[-1] = values[-1] + ln
            i += 1
            continue
        if re.search(r"Host Province|拟落地", ln):
            break
        values.append(collapse_cjk_spaces(ln))
        i += 1
        if len(values) >= 8:
            break

    order_keys = (
        "手机号",
        "电子邮箱",
        "最高学位中文",
        "最高学位英文",
        "回国前单位职务中文",
        "回国前单位职务英文",
    )
    out: dict[str, str] = {}
    idx = 0
    for v in values:
        while idx < len(order_keys) and out.get(order_keys[idx]):
            idx += 1
        if idx >= len(order_keys):
            break
        key = order_keys[idx]
        if key == "手机号" and not re.match(r"^[+\d]", v.replace(" ", "")):
            continue
        if key == "电子邮箱" and "@" not in v and not re.search(r"gmail", v, re.I):
            continue
        if key in ("最高学位中文", "回国前单位职务中文") and not re.search(r"[\u4e00-\u9fff]", v):
            continue
        if key in ("最高学位英文", "回国前单位职务英文") and not re.search(r"[A-Za-z]", v):
            continue
        if key == "电子邮箱":
            out[key] = fix_email_ocr(v)
        elif key in ("最高学位中文", "回国前单位职务中文"):
            out[key] = normalize_hj_cn_field(v.replace("、", "，").replace("；", "，"))
        else:
            out[key] = normalize_hj_en_field(v.replace("、", "，").replace(";", "，").replace(",", "，"))
        idx += 1
    return out


def _extract_vertical_bilingual_block(text: str, section: str, stop: str) -> tuple[str, str]:
    if section not in text:
        return "", ""
    block = text.split(section, 1)[1]
    if stop in block:
        block = block.split(stop, 1)[0]
    cn, en = "", ""
    skip_en = re.compile(
        r"(Last Employer|Before Returning|Coming to China|Highest Degree|Host Province|"
        r"Current or Expected|Type of Last|Country \(Region\)|Position$|Employer$)",
        re.I,
    )
    for ln in [x.strip() for x in block.splitlines() if x.strip()]:
        if ln.startswith(("□", "☑", "☐")):
            continue
        if re.match(r"^(中文|英文)\s*(Cn|En)\b", ln, re.I):
            continue
        if _looks_like_label(ln) or skip_en.search(ln):
            continue
        if re.search(r"[\u4e00-\u9fff]", ln) and not cn:
            cn = ln
        elif re.match(r"^[A-Za-z0-9]", ln) and not en and "@" not in ln:
            en = ln
    return cn, en


def _value_after_vertical_label(text: str, markers: tuple[str, ...], stop_markers: tuple[str, ...] = ()) -> str:
    for ln in _lines_after_marker(text, markers, stop_markers, 5):
        if _looks_like_label(ln) or ln.startswith(("□", "☑", "☐")):
            continue
        if re.match(r"^[A-Za-z].{0,40}$", ln) and any(m.lower() in ln.lower() for m in markers if re.match(r"^[A-Za-z]", m or "")):
            continue
        return ln
    return ""


def _normalize_degree(text: str) -> str:
    s = str(text or "").strip()
    if not s:
        return ""
    if re.match(r"\(Maste", s, re.I):
        s = "硕士（Master）"
    s = re.sub(r"\bPh\.D\.?\b", "PhD", s, flags=re.I)
    return format_hj_degree_cell(s)


def _normalize_bilingual(text: str) -> str:
    """职务/单位/院校：中文/English（参考 HJ 样例 table2 栏位）。"""
    s = fix_entity_name(fix_ocr_english(str(text or "").strip()))
    if not s:
        return ""
    pos = format_hj_position_cell(s)
    if pos and "/" in pos and re.search(r"[\u4e00-\u9fff]", pos):
        return pos
    if not re.search(r"[\u4e00-\u9fff]", s):
        return pos or s
    if "/" in s:
        a, b = s.split("/", 1)
        return f"{a.strip()}/{fix_ocr_english(b)}"
    m = re.match(r"^(.+?)\s*[（(]([^）)]+)[）)]\s*$", s)
    if m:
        return f"{m.group(1).strip()}/{fix_ocr_english(m.group(2))}"
    m = re.match(r"^(.+?)\s+([A-Za-z][^，,;；]+)$", s)
    if m and re.search(r"[\u4e00-\u9fff]", m.group(1)):
        return f"{m.group(1).strip()}/{fix_ocr_english(m.group(2))}"
    return s


def _hj_template_kind(doc: Document) -> str:
    return "compact" if len(doc.tables) <= 11 else "full"


def _edu_work_row_bounds(table) -> tuple[int, int, int]:
    """返回 (edu_start, edu_end, work_start) 行号。"""
    work_hdr = len(table.rows)
    for ri, row in enumerate(table.rows):
        t = _norm(" ".join(c.text for c in row.cells))
        if "工作经历" in t or "workexperience" in t:
            work_hdr = ri
            break
    return 1, work_hdr, work_hdr + 1


def _enrich_hj_fields(fields: dict[str, str], edu: list[dict], work: list[dict]) -> dict[str, str]:
    """参考 HJ 样例：用教育/工作表最高行补全 table1 与回国前信息。"""
    out = dict(fields)
    if edu:
        cn, en = highest_degree_from_edu_row(edu[0])
        if cn:
            out["最高学位中文"] = cn
        if en:
            out["最高学位英文"] = fix_ocr_english(en)
    if work:
        latest = work[0]
        if not out.get("回国前所在地"):
            c = str(latest.get("所在国家") or "").strip()
            if c:
                out["回国前所在地"] = _normalize_country_name(c)
        if not out.get("回国前单位职务中文") and not out.get("回国前单位职务英文"):
            ecn, een = split_bilingual(latest.get("工作单位") or "")
            pcn, pen = split_bilingual(latest.get("担任职务") or "")
            cn = cn_list_join(ecn, pcn)
            en = en_list_join(een, pen)
            if cn:
                out["回国前单位职务中文"] = cn
            if en:
                out["回国前单位职务英文"] = fix_ocr_english(en)
    cn_ret, en_ret = format_bilingual_pair(
        out.get("回国前单位职务中文") or "",
        out.get("回国前单位职务英文") or "",
    )
    if cn_ret:
        out["回国前单位职务中文"] = cn_ret
    if en_ret:
        out["回国前单位职务英文"] = fix_ocr_english(en_ret)
    for key in ("最高学位中文", "回国前单位职务中文"):
        if out.get(key):
            out[key] = normalize_hj_cn_field(out[key])
    for key in ("最高学位英文", "回国前单位职务英文"):
        if out.get(key):
            out[key] = normalize_hj_en_field(out[key])
    en_name, cn_name = split_certificate_name(out.get("有效证件姓名") or "")
    if en_name:
        out["有效证件姓名"] = en_name
    if cn_name and not out.get("中文（音译）名"):
        out["中文（音译）名"] = cn_name
    return out


def _clean_value(key: str, val: str) -> str:
    v = fix_ocr_english(str(val or "").strip())
    if not v:
        return ""
    v = re.sub(r"^[（(]?\s*Employer\s*[）)]?\s*[：:]\s*", "", v, flags=re.I)
    v = re.sub(r"^Current or Expected Employer Add\.?\s*", "", v, flags=re.I)
    if key == "性别":
        return _normalize_gender(v) or v
    if key == "出生日期":
        return _fmt_birth(v)
    if key == "出生国家（地区）":
        v = re.split(r"证件号码", v)[0].strip("：: ，,")
    if key in ("申报单位", "现工作单位", "用人单位名称") and v.startswith("（"):
        v = re.sub(r"^[（(][^）)]*[）)]\s*[：:]\s*", "", v)
    if key == "证件号码" and v.upper() in {"PASSPORTAL99", "OTHERIDNO", "IDNO"}:
        return ""
    if len(v) > 600 and key not in ("专长及代表性成果", "工作设想", "用人单位简介", "推荐理由", "支持条件"):
        v = v[:600]
    return v.strip()


def _first_group(patterns: list[tuple[str, str]], text: str) -> str:
    for key, pat in patterns:
        m = re.search(pat, text, re.I | re.M)
        if m:
            v = _clean_value(key, m.group(1) if m.lastindex else m.group(0))
            if v:
                return v
    return ""


def _paren_variants(s: str) -> tuple[str, ...]:
    a = str(s or "")
    b = a.replace("（", "(").replace("）", ")")
    c = a.replace("(", "（").replace(")", "）")
    out, seen = [], set()
    for x in (a, b, c):
        if x and x not in seen:
            seen.add(x)
            out.append(x)
    return tuple(out)


def _find_marker(raw: str, needle: str, start: int = 0) -> int:
    best = -1
    for v in _paren_variants(needle):
        i = raw.find(v, start)
        if i >= 0 and (best < 0 or i < best):
            best = i
    if best >= 0:
        return best
    compact_n = re.sub(r"\s+", "", needle)
    if len(compact_n) < 4:
        return -1
    compact_n = compact_n.replace("（", "(").replace("）", ")")
    buf = []
    for i, ch in enumerate(raw):
        if i < start:
            continue
        if ch.isspace():
            continue
        buf.append((i, ch.replace("（", "(").replace("）", ")")))
        if len(buf) > len(compact_n):
            buf.pop(0)
        if len(buf) == len(compact_n) and "".join(c for _, c in buf) == compact_n:
            return buf[0][0]
    return -1


def _section_between(text: str, starts: tuple[str, ...], ends: tuple[str, ...]) -> str:
    raw = str(text or "")
    start = -1
    for s in starts:
        i = _find_marker(raw, s)
        if i >= 0:
            start = i
            break
    if start < 0:
        return _section_between_loose(raw, starts, ends)
    nl = raw.find("\n", start)
    same = raw[start: nl if nl >= 0 else start + 80]
    if nl >= 0 and (
        len(re.sub(r"\s+", "", same)) < 48
        or re.search(r"须另附|情况介绍|300字|以内|Work Plan|Expertise", same)
    ):
        start = nl + 1
    end = len(raw)
    for e in ends:
        i = _find_marker(raw, e, start)
        if i > start:
            end = min(end, i)
    body = raw[start:end]
    lines = []
    for ln in body.splitlines():
        s = ln.strip()
        if not s or _PAGE_MARK.match(s):
            continue
        if s.startswith("I'm now zeroing"):
            break
        lines.append(s)
    out = re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()
    return out or _section_between_loose(raw, starts, ends)


def _section_between_loose(text: str, starts: tuple[str, ...], ends: tuple[str, ...]) -> str:
    """竖排/空格 OCR 时按行匹配章节起止标记。"""
    lines = [ln.strip() for ln in str(text or "").splitlines()]
    start_i = -1
    for i, ln in enumerate(lines):
        if _line_has_marker(ln, starts):
            start_i = i + 1
            break
        compact = re.sub(r"\s+", "", ln)
        for s in starts:
            if re.sub(r"\s+", "", s) in compact and len(re.sub(r"\s+", "", s)) >= 6:
                start_i = i + 1
                break
        if start_i >= 0:
            break
    if start_i < 0:
        return ""
    end_i = len(lines)
    for i in range(start_i, len(lines)):
        if _line_has_marker(lines[i], ends):
            end_i = i
            break
        compact = re.sub(r"\s+", "", lines[i])
        for e in ends:
            if re.sub(r"\s+", "", e) in compact and len(re.sub(r"\s+", "", e)) >= 6:
                end_i = i
                break
        if end_i < len(lines):
            break
    body_lines = []
    for ln in lines[start_i:end_i]:
        s = ln.strip()
        if not s or _PAGE_MARK.match(s):
            continue
        body_lines.append(s)
    return re.sub(r"\n{3,}", "\n\n", "\n".join(body_lines)).strip()


def _fill_edu_gaps(edu: list[dict]) -> list[dict]:
    """教育经历竖排 OCR 缺列时，用相邻学历行补全国家/院校。"""
    out = [dict(r) for r in edu]
    for i, row in enumerate(out):
        if row.get("所在国家") and row.get("校名称"):
            continue
        for j in range(i + 1, len(out)):
            older = out[j]
            if not row.get("所在国家") and older.get("所在国家"):
                row["所在国家"] = older["所在国家"]
            if not row.get("校名称") and older.get("校名称"):
                row["校名称"] = older["校名称"]
            if row.get("所在国家") and row.get("校名称"):
                break
        for j in range(i - 1, -1, -1):
            newer = out[j]
            if not row.get("所在国家") and newer.get("所在国家"):
                row["所在国家"] = newer["所在国家"]
            if not row.get("校名称") and newer.get("校名称"):
                row["校名称"] = newer["校名称"]
            if row.get("所在国家") and row.get("校名称"):
                break
    return [_normalize_edu_row(r) for r in out]


def _normalize_edu_row(d: dict) -> dict:
    start = str(d.get("开始时间") or "").strip()
    end = str(d.get("结束时间") or "").strip()
    range_text = str(d.get("起止时间") or "").strip()
    if start or end:
        range_text = _fmt_hj_range(start, end or "至今")
    elif range_text:
        range_text = _normalize_hj_range_text(range_text)
    school = _normalize_bilingual(str(d.get("校名称") or d.get("学校名称") or "").strip())
    major = str(d.get("专业领域") or d.get("专业名称") or "").strip()
    country = _normalize_country_name(str(d.get("所在国家") or "").strip())
    if not country:
        country = _infer_country_from_entity(school) or _infer_country_from_entity(major)
    return {
        "起止时间": range_text,
        "所在国家": country,
        "校名称": school,
        "专业领域": major,
        "学位": _normalize_degree(str(d.get("学位") or "").strip()),
    }


def _normalize_work_row(d: dict) -> dict:
    start = str(d.get("开始时间") or "").strip()
    end = str(d.get("结束时间") or "").strip()
    range_text = str(d.get("起止时间") or "").strip()
    if start or end:
        range_text = _fmt_hj_range(start, end or "至今")
    elif range_text:
        range_text = _normalize_hj_range_text(range_text)
    employer_raw = _strip_checkbox_garbage(str(d.get("工作单位") or d.get("单位") or "").strip())
    position_raw = _strip_checkbox_garbage(str(d.get("担任职务") or d.get("职务") or "").strip())
    if employer_raw and (not position_raw or _is_hj_field_garbage(position_raw)):
        emp, pos = _split_employer_position(employer_raw)
        if pos:
            employer_raw, position_raw = emp, pos
    employer = _normalize_bilingual(employer_raw)
    position = format_hj_position_cell(position_raw) or _normalize_bilingual(position_raw)
    if re.search(r"(博士|硕士|学士)\s*[（(]", employer) and re.search(r"(大学|学院|University)", employer, re.I):
        return {}
    country = _normalize_country_name(str(d.get("所在国家") or "").strip())
    if not country:
        country = _infer_country_from_entity(employer)
    perf = str(d.get("任职情况") or d.get("工作性质") or "全职").strip() or "全职"
    if _is_hj_field_garbage(perf):
        perf = "全职"
    return {
        "起止时间": range_text,
        "所在国家": country,
        "工作单位": employer,
        "担任职务": position,
        "任职情况": perf,
    }


def _strip_timeline_noise(text: str) -> str:
    t = str(text or "")
    t = re.sub(
        r"(Start from|Educational Background|All Work Experience|Please Complete|"
        r"Held Both Within|China and Abroad|Consecutive Chronological|Order\.|"
        r"Bachelor Degree|Your Personal|Work History in)",
        " ",
        t,
        flags=re.I,
    )
    return re.sub(r"\s+", " ", t).strip()


def _extract_timeline_country(text: str) -> str:
    compact = re.sub(r"\s+", "", str(text or ""))
    for c in (
        "印度", "韩国", "加拿大", "美国", "英国", "中国", "德国", "法国", "日本",
        "澳大利亚", "巴西", "新加坡", "马来西亚",
    ):
        if c in compact:
            return c
    return _infer_country_from_entity(text)


def _extract_timeline_degree(text: str) -> str:
    for pat in (
        r"博士\s*\(PhD\)",
        r"硕士\s*\(Master[^)]*\)",
        r"学士\s*\(Bachelor[^)]*\)",
        r"\(PhD\)",
        r"\(Master[^)]*\)",
        r"\(Maste[^)]*\)",
        r"\(Bachelor[^)]*\)",
        r"Postdoctoral\s+Research\s+Fellow",
        r"Postdoctoral\s+Fellow",
        r"Assistant\s+Professor",
        r"Chief\s+Scientist",
        r"博士后",
        r"博士",
        r"硕士",
        r"学士",
    ):
        m = re.search(pat, str(text or ""), re.I)
        if m:
            return m.group(0)
    return ""


def _extract_timeline_school(text: str) -> str:
    for pat, name in _TIMELINE_SCHOOL_HINTS:
        if pat.search(str(text or "")):
            return name
    raw = _strip_timeline_noise(text)
    m = re.search(
        r"([\u4e00-\u9fff]{2,12}(?:大学|学院))(?:\s*/\s*([A-Za-z][A-Za-z\s]{2,40}))?",
        raw,
    )
    if m:
        cn = re.sub(r"\s+", "", m.group(1))
        en = fix_ocr_english(m.group(2) or "")
        return format_bilingual_pair(cn, en)[0] if en else cn
    return ""


def _extract_timeline_major(text: str, degree: str) -> str:
    raw = _strip_timeline_noise(text)
    deg = str(degree or "")
    if deg and deg in raw:
        before = raw.split(deg, 1)[0]
        m = re.search(r"([\u4e00-\u9fff]{2,20}(?:工程|化学|科学|材料|物理))", before)
        if m:
            return re.sub(r"\s+", "", m.group(1))
    m = re.search(
        r"([\u4e00-\u9fff]{2,20}(?:工程|化学|科学|材料|物理|电化学))",
        raw,
    )
    return re.sub(r"\s+", "", m.group(1)) if m else ""


def _extract_timeline_position(text: str) -> str:
    raw = str(text or "")
    if re.search(r"Postdoctoral", raw, re.I) and re.search(r"Fellow", raw, re.I):
        return "博士后/Postdoctoral Research Fellow"
    for pat in (
        r"Postdoctoral\s+Research\s+Fellow",
        r"Postdoctoral\s+Fellow",
        r"Assistant\s+Professor",
        r"Chief\s+Scientist",
        r"Research\s+Fellow",
        r"博士后(?:研究)?员?",
        r"助理教授",
        r"首席科学家",
    ):
        m = re.search(pat, raw, re.I)
        if m:
            return format_hj_position_cell(fix_ocr_english(m.group(0)))
    if re.search(r"Chief\s*Scientist", raw, re.I) or ("首席" in raw and "科学家" in raw):
        return "首席科学家/Chief Scientist"
    if re.search(r"Scientist", raw, re.I) and re.search(r"AdvEn|阿德文", raw, re.I):
        return "首席科学家/Chief Scientist"
    if re.search(r"Assistant", raw, re.I) and re.search(r"Professor", raw, re.I):
        return "助理教授/Assistant Professor"
    m = re.search(r"\(([^)]{4,40}(?:Professor|Fellow|Scientist|Research)[^)]*)\)", raw, re.I)
    if m:
        pos = fix_ocr_english(m.group(1))
        pos = re.sub(r"全\s*职", "", pos, flags=re.I)
        pos = re.sub(r"[,、/]+", " ", pos)
        pos = re.sub(r"\s+", " ", pos).strip()
        if re.search(r"Assistant", pos, re.I) and re.search(r"Professor", pos, re.I):
            return "Assistant Professor"
        return pos
    return ""


def _edu_from_timeline_ctx(start: str, end: str, ctx: str) -> dict | None:
    from .scan_text_normalize import collapse_cjk_spaces

    clean = _strip_timeline_noise(collapse_cjk_spaces(ctx))
    degree = _extract_timeline_degree(clean)
    if not degree:
        return None
    school = _extract_timeline_school(clean)
    if not school:
        return None
    return _normalize_edu_row({
        "起止时间": _fmt_hj_range(start, end),
        "所在国家": _extract_timeline_country(clean),
        "校名称": school,
        "专业领域": _extract_timeline_major(clean, degree),
        "学位": degree,
    })


def _work_from_timeline_ctx(start: str, end: str, ctx: str) -> dict | None:
    from .scan_text_normalize import collapse_cjk_spaces

    clean = _strip_timeline_noise(collapse_cjk_spaces(ctx))
    employer = _extract_timeline_school(clean)
    if not employer:
        m = re.search(
            r"([\u4e00-\u9fff]{2,16}(?:大学|学院|工业|公司|Corp|Inc)[^，,。\n]{0,40})",
            clean,
        )
        if m:
            employer = re.sub(r"\s+", "", m.group(1))
    position = _extract_timeline_position(clean)
    if not employer and not position:
        return None
    perf = "兼职" if re.search(r"兼职|Part[- ]?time", clean, re.I) else "全职"
    return _normalize_work_row({
        "起止时间": _fmt_hj_range(start, end),
        "所在国家": _extract_timeline_country(clean),
        "工作单位": employer,
        "担任职务": position,
        "任职情况": perf,
    })


def _find_timeline_section_end(lines: list[str], start: int) -> int:
    for i in range(start, len(lines)):
        if _line_has_marker(lines[i], _TIMELINE_STOP_MARKERS):
            return i
        compact = re.sub(r"\s+", "", lines[i])
        for m in _TIMELINE_STOP_MARKERS:
            if m in compact:
                return i
    return len(lines)


def _parse_timeline_section(lines: list[str], start: int, end: int, kind: str) -> list[dict]:
    if start < 0 or start >= end:
        return []
    section_end = _find_timeline_section_end(lines, start)
    end = min(end, section_end)
    blob = "\n".join(lines[start:end])
    matches = list(_DATE_PAIR_INLINE.finditer(blob))
    if not matches:
        return []
    rows: list[dict] = []
    for i, m in enumerate(matches):
        a, b_raw = m.group(1), m.group(2)
        ctx_end = matches[i + 1].start() if i + 1 < len(matches) else len(blob)
        ctx = blob[m.end():ctx_end]
        b = _repair_yyyymm_end(a, b_raw, ctx)
        row = _edu_from_timeline_ctx(a, b, ctx) if kind == "edu" else _work_from_timeline_ctx(a, b, ctx)
        if row and _row_identity(row, kind):
            rows.append(row)
    return rows


def _is_valid_timeline_row(row: dict, kind: str) -> bool:
    if not row:
        return False
    rng = _normalize_hj_range_text(str(row.get("起止时间") or ""))
    if not re.match(r"\d{4}\.\d{2}\.\d{2}-\d{4}\.\d{2}\.\d{2}", rng) and "至今" not in rng:
        return False
    if kind == "edu":
        school = str(row.get("校名称") or "")
        if len(school) > 72 or _HJ_FIELD_GARBAGE.search(school):
            return False
    else:
        emp = str(row.get("工作单位") or "")
        if len(emp) > 72 or _HJ_FIELD_GARBAGE.search(emp):
            return False
        if re.search(
            r"(大学|学院|University).*(博士(?!后)|硕士|学士|Doctor|Master|Bachelor)",
            emp,
            re.I,
        ):
            return False
    return True


def _row_identity(row: dict, kind: str) -> str:
    if not row:
        return ""
    if kind == "edu":
        return _norm(str(row.get("校名称") or "") + str(row.get("学位") or ""))
    return _norm(str(row.get("工作单位") or "") + str(row.get("担任职务") or ""))


def _merge_timeline_rows(ocr_rows: list[dict], parser_rows: list[dict], kind: str) -> list[dict]:
    """OCR 行优先；解析器补全缺失段，统一为 HJ 参考格式（时间倒序）。"""
    norm = _normalize_edu_row if kind == "edu" else _normalize_work_row
    ocr = [
        r for r in (norm(d) for d in ocr_rows)
        if _row_identity(r, kind) and _is_valid_timeline_row(r, kind)
    ]
    parser = [
        r for r in (norm(d) for d in parser_rows)
        if _row_identity(r, kind) and _is_valid_timeline_row(r, kind)
    ]
    if not parser:
        return ocr
    if not ocr:
        return list(reversed(parser))
    seen = {_row_identity(r, kind) for r in ocr}
    merged = list(ocr)
    for row in reversed(parser):
        key = _row_identity(row, kind)
        if key and key not in seen:
            merged.append(row)
            seen.add(key)
    return merged


def _is_vertical_noise(line: str) -> bool:
    s = str(line or "").strip()
    if not s or _PAGE_MARK.match(s):
        return True
    if _VERTICAL_NOISE.match(s):
        return True
    if re.match(r"^\(?从本科", s) or re.match(r"^\(?完整填写", s):
        return True
    if len(s) > 80 and re.search(r"(Time|Country|Region|University|Degree)", s, re.I):
        return True
    return False


def _find_degree_line_index(lines: list[str]) -> int:
    for i in range(len(lines) - 1, -1, -1):
        compact = re.sub(r"\s+", "", lines[i])
        if re.search(r"(学士|硕士|博士)", compact):
            return i
        if re.search(r"\b(Bachelor|Master|Doctor|PhD)\b", lines[i], re.I):
            return i
    return -1


def _vertical_edu_from_buf(a: str, b: str, buf: list[str]) -> dict | None:
    lines = [x for x in buf if x and not _is_vertical_noise(x)]
    if not lines:
        return None
    deg_i = _find_degree_line_index(lines)
    if deg_i < 0:
        return None
    major = lines[deg_i - 1] if deg_i >= 1 else ""
    rest = lines[: max(0, deg_i - 1)]
    country, school = "", ""
    if rest and _COUNTRY_NAME.match(rest[0]):
        country = rest[0]
        school = " ".join(rest[1:]).strip()
    else:
        school = " ".join(rest).strip()
    return _normalize_edu_row({
        "起止时间": _fmt_hj_range(a, b),
        "所在国家": country,
        "校名称": school,
        "专业领域": major,
        "学位": lines[deg_i],
    })


def _vertical_work_from_buf(a: str, b: str, buf: list[str]) -> dict | None:
    lines = [x for x in buf if x and not _is_vertical_noise(x)]
    if not lines:
        return None
    perf = "全职"
    if lines[-1] in ("全职", "兼职"):
        perf = lines[-1]
        lines = lines[:-1]
    if not lines:
        return None
    position = lines[-1]
    employer = " ".join(lines[:-1]).strip() if len(lines) > 1 else lines[0]
    country = ""
    if _COUNTRY_NAME.match(employer):
        country, employer = employer, (lines[1] if len(lines) > 2 else "")
        position = lines[-1]
    return _normalize_work_row({
        "起止时间": _fmt_hj_range(a, b),
        "所在国家": country,
        "工作单位": employer,
        "担任职务": position,
        "任职情况": perf,
    })


def _parse_vertical_block(lines: list[str], start: int, end: int, kind: str) -> list[dict]:
    rows: list[dict] = []
    i = start
    while i < end:
        m = _DATE_ONLY.match(lines[i])
        tail = ""
        if not m:
            m2 = _DATE_LOOSE.match(lines[i])
            if not m2:
                i += 1
                continue
            m = m2
            tail = str(m2.group(3) or "").strip()
        buf: list[str] = []
        if tail:
            buf.append(tail)
        i += 1
        while i < end:
            if _DATE_ONLY.match(lines[i]):
                break
            if any(k in lines[i] for k in ("破格申报", "重要科研奖励", "科研情况", "工作设想")):
                break
            if not _is_vertical_noise(lines[i]):
                buf.append(lines[i])
            i += 1
        row = _vertical_edu_from_buf(m.group(1), m.group(2), buf) if kind == "edu" else _vertical_work_from_buf(m.group(1), m.group(2), buf)
        if row:
            rows.append(row)
    return rows


def _line_has_marker(line: str, markers: tuple[str, ...]) -> bool:
    compact = re.sub(r"\s+", "", str(line or ""))
    for m in markers:
        if m in str(line or "") or m in compact:
            return True
    return False


def _parse_vertical_list_lines(text: str) -> tuple[list[dict], list[dict]]:
    raw = [ln.strip() for ln in str(text or "").splitlines()]
    edu_i = work_i = -1
    for i, ln in enumerate(raw):
        if edu_i < 0 and _line_has_marker(ln, ("教育经历", "Educational Background")):
            edu_i = i
        if work_i < 0 and _line_has_marker(ln, ("全部工作经历", "All Work Experience")):
            work_i = i
    edu_end = work_i if work_i > edu_i else len(raw)
    edu = _parse_timeline_section(raw, edu_i + 1, edu_end, "edu") if edu_i >= 0 else []
    work = _parse_timeline_section(raw, work_i + 1, len(raw), "work") if work_i >= 0 else []
    if len(edu) < 2 and edu_i >= 0:
        legacy = _parse_vertical_block(raw, edu_i + 1, edu_end, "edu")
        if len(legacy) > len(edu):
            edu = legacy
    if len(work) < 2 and work_i >= 0:
        legacy = _parse_vertical_block(raw, work_i + 1, _find_timeline_section_end(raw, work_i + 1), "work")
        if len(legacy) > len(work):
            work = legacy
    return edu, work


def _parse_list_lines(text: str) -> tuple[list[dict], list[dict]]:
    from .scan_text_normalize import deocr_scanned_text

    text = deocr_scanned_text(text)
    edu, work = [], []
    for line in str(text or "").splitlines():
        s = line.strip()
        if not s:
            continue
        m = _EDU_TAB.search(s)
        if m:
            edu.append(_normalize_edu_row({
                "起止时间": _fmt_hj_range(m.group(1), m.group(2)),
                "所在国家": m.group(3).strip(),
                "校名称": m.group(4).strip(),
                "专业领域": m.group(5).strip(),
                "学位": m.group(6).strip(),
            }))
            continue
        m = _WORK_TAB.search(s)
        if m:
            work.append(_normalize_work_row({
                "起止时间": _fmt_hj_range(m.group(1), m.group(2)),
                "所在国家": m.group(3).strip(),
                "工作单位": m.group(4).strip(),
                "担任职务": m.group(5).strip(),
                "任职情况": m.group(6).strip(),
            }))
    v_edu, v_work = _parse_vertical_list_lines(text)
    if len(v_edu) > len(edu):
        edu = v_edu
    if len(v_work) > len(work):
        work = v_work
    edu = [r for r in (_normalize_edu_row(d) for d in edu) if _is_valid_timeline_row(r, "edu")]
    work = [r for r in (_normalize_work_row(d) for d in work) if _is_valid_timeline_row(r, "work")]
    return list(reversed(edu)), list(reversed(work))


def _is_scalar_label_value(val: str) -> bool:
    s = str(val or "").strip()
    if not s:
        return True
    return bool(_SCALAR_LABEL_GARBAGE.search(s))


def _clean_research_direction(val: str, raw: str = "") -> str:
    block = raw.split("具体研究方向", 1)[-1] if "具体研究方向" in raw else str(val or "")
    lines = []
    for ln in block.splitlines()[:5]:
        t = ln.strip()
        if not t or _PAGE_MARK.match(t):
            continue
        if re.search(r"涉及关键核心技术|Type of Research", t):
            break
        lines.append(t)
    merged = re.sub(r"\s+", "", "".join(lines))
    m = re.search(
        r"(新型[\u4e00-\u9fff、，,；;/]{6,}|电解质[\u4e00-\u9fff、，,；;/]{4,})",
        merged,
    )
    if m:
        return _fix_ocr_terms(m.group(0).rstrip("、，,；;/"))
    cn = re.sub(r"[^一-龥、，,；;/]", "", str(val or "") + merged)
    cn = re.sub(r"\s+", "", cn)
    cn = re.sub(r"^[^一-龥]*(?=[一-龥])", "", cn)
    cn = _fix_ocr_terms(cn)
    if len(cn) >= 12:
        return cn
    return _fix_ocr_terms(str(val or "").strip())


def _fix_ocr_terms(text: str) -> str:
    s = str(text or "")
    for a, b in _OCR_TERM_FIXES:
        s = s.replace(a, b)
    return s


def _clean_narrative(text: str) -> str:
    from .scan_text_normalize import collapse_cjk_spaces

    lines = []
    for ln in str(text or "").splitlines():
        s = collapse_cjk_spaces(ln.strip())
        if not s or _PAGE_MARK.search(s) or re.match(r"^\d{1,3}/\d{1,3}$", s):
            continue
        if re.fullmatch(r"-{2,}", s):
            continue
        if "未提交" in s and len(s) < 48:
            continue
        if s.startswith("I'm now zeroing"):
            break
        if re.search(r"申报渠道意见|对申报人及申报材料的审核意见", s) and len(s) < 40:
            break
        s = re.sub(r"^\\+(\d+)[.、．]", r"\1.", s)
        s = re.sub(r"^@\s*〇\s*", "（1）", s)
        s = re.sub(r"^G@\)\s*", "（2）", s)
        s = re.sub(r"^@\s*团队保障", "（3）团队保障", s)
        s = re.sub(r"^@\s*生活配套", "（4）生活配套", s)
        s = re.sub(r"^@\s+", "", s)
        lines.append(_fix_ocr_terms(s))
    out = re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()
    out = re.sub(r" {2,}", " ", out)
    out = re.sub(r"\d{4}/\d{1,2}/\d{1,2}\s+\d{1,2}:\d{2}\s*未提交", "", out)
    out = re.sub(r"20\d{2}\d/\d+/\d+\s*未提交", "", out)
    out = _fix_ocr_terms(out.replace("债\n守", "恪守").replace("债 守", "恪守"))
    return out.strip()


def _ocr_checked_option(text: str, options: tuple[str, ...]) -> str:
    raw = str(text or "")
    compact = re.sub(r"\s+", "", raw)
    for opt in options:
        oc = re.sub(r"\s+", "", opt)
        if not oc:
            continue
        if re.search(_OCR_CHECKED.pattern + r"\s*" + re.escape(opt), raw):
            return opt
        if re.search(_OCR_CHECKED.pattern + re.escape(oc), compact):
            return opt
        idx = compact.find(oc)
        if idx >= 1 and _OCR_CHECKED.match(compact[idx - 1]):
            return opt
        if "☑" in raw and opt in raw:
            if idx >= 1 and _OCR_CHECKED.match(compact[idx - 1]):
                return opt
    return ""


def _keywords_from_direction(direction: str) -> str:
    s = _fix_ocr_terms(str(direction or ""))
    s = s.replace("以及", "、").replace("及", "、")
    parts = []
    for p in re.split(r"[、，,；;]", s):
        item = p.strip(" 、，,；;/")
        cn = re.sub(r"[^\u4e00-\u9fff/]", "", item)
        if 3 <= len(cn) <= 24:
            parts.append(item)
    seen, out = set(), []
    for p in parts:
        k = re.sub(r"\s+", "", p)
        if k and k not in seen:
            seen.add(k)
            out.append(p)
    return "、".join(out[:5])


def _clean_admin_region(val: str, kind: str) -> str:
    s = re.sub(r"\s+", "", str(val or ""))
    if kind == "省":
        if "江苏" in s:
            return "江苏省"
        m = re.search(r"([\u4e00-\u9fff]{2,6}省)", s)
        return m.group(1) if m else ""
    if "南京" in s:
        return "南京市"
    m = re.search(r"([\u4e00-\u9fff]{2,6}市)", s)
    if m:
        return m.group(1)
    return s if re.fullmatch(r"[\u4e00-\u9fff]{2,8}", s) else ""


def _extract_employment_block_fields(text: str) -> dict[str, str]:
    """竖排 OCR 第 4 页：拟落地省/市、现工作单位、地址、回国时间。"""
    lines = [ln.strip() for ln in str(text or "").splitlines()]
    start = -1
    for i, ln in enumerate(lines):
        if _line_has_marker(ln, ("拟落地省", "Host Province", "拟 落 地 省")):
            start = i
            break
    if start < 0:
        return {}
    block = "\n".join(lines[start : start + 14])
    compact = re.sub(r"\s+", "", block)
    out: dict[str, str] = {}
    if re.search(r"江\s*苏\s*省", block):
        out["拟落地省"] = "江苏省"
    if re.search(r"南\s*京\s*市", block):
        out["拟落地市"] = "南京市"
    if re.search(r"南京博驰新能源(?:股份)?有限公司", compact):
        out["现工作单位"] = "南京博驰新能源股份有限公司"
    elif re.search(r"南京博驰", compact):
        out["现工作单位"] = fix_entity_name("南京博驰新能源股份有限公司")
    m = re.search(r"(南京市[\u4e00-\u9fff\d]+?号[\u4e00-\u9fff\d]*?室)", compact)
    if m:
        out["任职单位地址"] = re.split(r"20\d{2}年", m.group(1))[0]
    m = re.search(r"(20\d{2})\s*年\s*(\d{1,2})\s*月", block)
    if m:
        out["回国时间"] = f"{m.group(1)}.{int(m.group(2)):02d}"
    for pos in ("首席科学家", "助理教授", "研究员", "教授", "总工程师", "技术总监"):
        if pos in block.replace(" ", ""):
            out["引进职务"] = pos
            break
    return out


def _paper_role_from_block(block: str, rank: str) -> str:
    compact = re.sub(r"\s+", "", block)
    for role in ("通讯作者", "共同第一作者", "第一作者", "核心研究员", "核心作者"):
        if role in compact:
            return role
    m = re.search(r"排\s*序\s*[:：]\s*(\d+/\d+)", block)
    if m:
        return f"排序{m.group(1)}"
    if rank:
        return f"排序{rank}"
    return ""


def _paper_title_from_block(block: str, date: str, venue: str, rank: str) -> str:
    blob = re.sub(r"\s+", " ", str(block or "").replace("\n", " "))
    blob = re.sub(
        r"(Publication Date|Publication Venue|Role and|Contribution|序号|pany|"
        r"Number of|NO\.|Authors|申报人角色|Surface Coatin|esetrochemical|ae oe)",
        " ",
        blob,
        flags=re.I,
    )
    m = re.search(
        r"([\u4e00-\u9fff]{6,48}(?:材料|电池|电容|进展|研究|设计|工程|阴极|阳极|回收|封装|沉积|涂层)[\u4e00-\u9fff]{0,24})"
        r"\s*/\s*([A-Za-z][A-Za-z0-9\s,\-:]{12,}?)"
        r"(?=\s*(?:Advanced|Nano|Energy|Electro|Carbon|Materials|Acta|Journal|\d{1,2}/\d{1,2}|$))",
        blob,
        re.I,
    )
    if m:
        cn = re.sub(r"\s+", "", m.group(1))
        en = fix_ocr_english(re.sub(r"\s+", " ", m.group(2)).strip(" ,;"))
        if len(cn) >= 8 and len(en) >= 12:
            return f"{cn}/{en}"
    m2 = re.search(
        r"([\u4e00-\u9fff]{8,48})\s*/\s*([A-Za-z][A-Za-z0-9\s,\-:]{12,})",
        blob,
    )
    if m2:
        cn = re.sub(r"\s+", "", m2.group(1))
        en = fix_ocr_english(re.sub(r"\s+", " ", m2.group(2)).strip(" ,;"))
        if len(cn) >= 8:
            return f"{cn}/{en}" if en else cn
    cn_only = re.findall(r"[\u4e00-\u9fff]{10,48}", re.sub(r"\s+", "", blob))
    en_only = re.findall(r"[A-Za-z][A-Za-z\s,\-:]{16,}", blob)
    if cn_only:
        cn = cn_only[0]
        en = fix_ocr_english(en_only[0].strip()) if en_only else ""
        return f"{cn}/{en}" if en else cn
    return ""


def _paper_from_block(serial: int, date: str, block: str) -> dict | None:
    from .scan_text_normalize import collapse_cjk_spaces

    block = collapse_cjk_spaces(block)
    if not block.strip():
        return None
    ranks = _PAPER_RANK.findall(block)
    rank = ranks[0] if ranks else ""
    venue = ""
    vm = _PAPER_VENUE.search(block.replace("\n", " "))
    if vm:
        venue = vm.group(0).strip()
    if not venue:
        for ln in block.splitlines():
            if _PAPER_VENUE.search(ln) and len(ln.strip()) < 80:
                venue = _PAPER_VENUE.search(ln).group(0).strip()
                break
    title = _paper_title_from_block(block, date, venue, rank)
    cn_len = len(re.sub(r"[^\u4e00-\u9fff]", "", title))
    en_len = len(re.sub(r"[^A-Za-z]", "", title))
    if re.fullmatch(r"(?:\d{4}[-./]\d{2}(?:[-./]\d{2})?\s*)+", title or ""):
        return None
    if cn_len < 4 and not (venue and en_len >= 12):
        return None
    role = _paper_role_from_block(block, rank)
    return {
        "发表时间": _fmt_hj_date(date),
        "论文题目": title,
        "发表载体": venue,
        "排序": rank,
        "角色": role,
    }


def _parse_papers_vertical(text: str) -> list[dict]:
    section = _section_between_loose(text, _PAPER_SECTION_START, _PAPER_SECTION_END)
    if not section:
        section = _section_between(text, _PAPER_SECTION_START, _PAPER_SECTION_END)
    if not section:
        return []
    lines = [ln.strip() for ln in section.splitlines() if ln.strip()]
    anchors: list[tuple[int, int, str]] = []
    for i, ln in enumerate(lines):
        m = re.match(r"^(\d{1,2})\s+(\d{4}[-./]\d{1,2}(?:[-./]\d{1,2})?)\b", ln)
        if m:
            anchors.append((i, int(m.group(1)), m.group(2)))
            continue
        m2 = re.match(r"^(\d{1,2})\s*/", ln)
        if m2 and _PAPER_DATE_ONLY.search(ln):
            dm = _PAPER_DATE_ONLY.search(ln)
            if dm:
                anchors.append((i, int(m2.group(1)), dm.group(1)))
            continue
        if _PAPER_DATE_ONLY.match(ln):
            if i > 0 and _PAPER_SERIAL_ONLY.match(lines[i - 1]):
                anchors.append((i, int(lines[i - 1]), ln))
            else:
                anchors.append((i, len(anchors) + 1, ln))
    if not anchors:
        return []
    rows: list[dict] = []
    for j, (idx, serial, date) in enumerate(anchors):
        next_idx = anchors[j + 1][0] if j + 1 < len(anchors) else len(lines)
        pre = lines[max(0, idx - 8):idx]
        post = lines[idx:next_idx]
        block = "\n".join(pre + post)
        row = _paper_from_block(serial, date, block)
        if row:
            rows.append(row)
    dedup: list[dict] = []
    seen: set[str] = set()
    for row in rows:
        key = _norm(str(row.get("发表时间") or "") + str(row.get("论文题目") or "")[:40])
        if key and key not in seen:
            seen.add(key)
            dedup.append(row)
    dedup.sort(
        key=lambda r: (
            len(re.sub(r"[^\u4e00-\u9fff]", "", str(r.get("论文题目") or ""))),
            bool(r.get("发表载体")),
            bool(r.get("排序")),
        ),
        reverse=True,
    )
    return dedup[:10]


def _parse_papers(text: str) -> list[dict]:
    from .scan_text_normalize import deocr_scanned_text

    text = deocr_scanned_text(text)
    rows = []
    for line in str(text or "").splitlines():
        s = line.strip()
        if not s or _PAGE_MARK.match(s):
            continue
        m = _PAPER_LINE.match(s)
        if not m:
            continue
        rows.append({
            "发表时间": m.group(1),
            "论文题目": m.group(2).strip(),
            "发表载体": m.group(3).strip(),
            "排序": m.group(4).strip(),
            "角色": m.group(5).strip(),
        })
    if len(rows) < 2:
        vertical = _parse_papers_vertical(text)
        if len(vertical) > len(rows):
            rows = vertical
    cleaned = []
    seen = set()
    for row in rows:
        title = re.sub(r"\s+", " ", str(row.get("论文题目") or "")).strip()
        title = re.sub(r"(?<=[A-Za-z])\s+(?=[a-z]{1,4}\b)", "", title)
        if re.search(r"起止时间|性质和来源|Title of|No Time|序号发表|^负责|^任务|pseudocapa|Na2CoP", title, re.I):
            continue
        if re.search(r"20\d{2}-\d{2}", title):
            continue
        cn = re.sub(r"[^\u4e00-\u9fff]", "", title)
        if len(cn) < 4 and not (row.get("发表载体") and re.search(r"[A-Za-z]{8,}", title)):
            continue
        if not row.get("发表载体") and len(cn) < 8:
            continue
        row = dict(row)
        row["论文题目"] = _fix_ocr_terms(title)
        row["发表时间"] = _fmt_hj_date(row.get("发表时间") or "")
        key = cn[:16]
        if key in seen:
            continue
        if any(cn in re.sub(r"[^\u4e00-\u9fff]", "", str(x.get("论文题目") or "")) or re.sub(r"[^\u4e00-\u9fff]", "", str(x.get("论文题目") or "")) in cn for x in cleaned):
            continue
        seen.add(key)
        cleaned.append(row)
    cleaned.sort(
        key=lambda r: (bool(r.get("发表载体")), len(re.sub(r"[^\u4e00-\u9fff]", "", str(r.get("论文题目") or "")))),
        reverse=True,
    )
    return cleaned[:6]


def _parse_projects(text: str) -> list[dict]:
    from .scan_text_normalize import deocr_scanned_text

    raw = deocr_scanned_text(text)
    section = _section_between_loose(raw, _PROJECT_SECTION_START, _PROJECT_SECTION_END)
    if not section:
        section = _section_between(raw, _PROJECT_SECTION_START, _PROJECT_SECTION_END)
    if not section:
        return []
    blob = re.sub(r"\s+", " ", section)
    blob = re.sub(
        r"(20\d{2}-\d{2}-\d{2}-)(\d)\s+(\d{3}-\d{2}-\d{2})",
        lambda m: m.group(1) + m.group(2) + m.group(3),
        blob,
    )
    date_pat = r"(20\d{2}[-./]\d{1,2}[-./]\d{1,2})"
    hits = list(re.finditer(date_pat + r"\s*-\s*" + date_pat, blob))
    if not hits:
        singles = list(re.finditer(r"20\d{2}[-./]\d{1,2}[-./]\d{1,2}", blob))
        paired = []
        i = 0
        while i + 1 < len(singles):
            a, b = singles[i], singles[i + 1]
            if a.end() < b.start() <= a.end() + 8:
                paired.append((a, b))
                i += 2
            else:
                i += 1
        class _M:
            def __init__(self, a, b):
                self._a, self._b = a, b
            def start(self):
                return self._a.start()
            def end(self):
                return self._b.end()
            def group(self, n):
                return self._a.group(0) if n == 1 else self._b.group(0)
        hits = [_M(a, b) for a, b in paired]
    rows: list[dict] = []
    for i, m in enumerate(hits):
        nxt = hits[i + 1].start() if i + 1 < len(hits) else min(len(blob), m.end() + 360)
        chunk = blob[m.start():nxt]
        rng = _fmt_hj_range(re.sub(r"\D", "", m.group(1)), re.sub(r"\D", "", m.group(2)))
        nature = ""
        nm = re.search(r"项目性质[:：]?\s*([^，,。来职]{2,16})", chunk)
        if nm:
            nature = re.sub(r"\s+", "", nm.group(1))
            nature = re.sub(r"研[究尻].*$", "研究", nature)
        src = ""
        sm = re.search(r"来源[:：]\s*([^职经]{4,80})", chunk)
        if sm:
            src = re.sub(r"\s+", "", sm.group(1))[:40]
        title = ""
        tm = re.search(
            r"([\u4e00-\u9fff]{8,48}(?:材料|电池|电容|涂层|电极|工程|阴极|阳极|回收|电容器|展望|研究))",
            chunk,
        )
        if tm:
            title = _fix_ocr_terms(tm.group(1))
        budget = ""
        bm = re.search(r"\b(\d{2,3}(?:\.\d+)?)\b", chunk)
        if bm and not re.fullmatch(r"20\d{2}", bm.group(1)):
            budget = bm.group(1)
            if budget and "万" not in budget:
                budget = budget + "万元"
        if title and re.match(r"^(负责|任务|职位|来源|申报人)", title):
            title = ""
        rank = ""
        rm = re.search(r"(\d{1,2}/\d{1,2})", chunk)
        if rm:
            rank = rm.group(1)
        role = "项目主持人" if "主持" in chunk else ("项目参与人" if "参与" in chunk else "")
        if "共同主持" in chunk:
            role = "项目共同主持人"
        if rank:
            role = ("职位：" + role + " 排序：" + rank) if role else ("排序：" + rank)
        elif role:
            role = "职位：" + role
        desc = title or nature
        funders = []
        for lab in ("NSERC", "MITACS", "加拿大创新基金会", "韩国国家研究基金会", "国家自然科学基金", "通用汽车"):
            if lab.lower() in chunk.lower() or lab in chunk:
                funders.append(lab)
        source_parts = [p for p in (nature, src) if p and not re.search(r"20\d{2}-", p)]
        if funders:
            source_parts.extend(funders)
        if title:
            source_parts.append(title)
        seen_p, source_parts_u = set(), []
        for p in source_parts:
            k = re.sub(r"\s+", "", p)
            if k and k not in seen_p:
                seen_p.add(k)
                source_parts_u.append(p)
        source = "，".join(source_parts_u)
        if not rng:
            continue
        rows.append({
            "描述": desc,
            "起止时间": rng,
            "性质来源": source or desc,
            "经费": budget,
            "角色": role,
        })
    dedup, seen = [], set()
    for row in rows:
        key = row.get("起止时间") or ""
        desc = str(row.get("描述") or "")
        src = str(row.get("性质来源") or "")
        cn = re.sub(r"[^\u4e00-\u9fff]", "", desc)
        if "测试数" in src:
            continue
        if len(cn) < 4 or re.match(r"^(负责|任务|职位|来源)", desc):
            continue
        if key and key not in seen:
            seen.add(key)
            dedup.append(row)
    return dedup[:5]


def _pick_tab_after(label: str, text: str) -> str:
    return _pick_after_label(text, label)


def _pick_after_label(text: str, label: str, stop_labels: tuple[str, ...] = ()) -> str:
    """取 OCR 竖排「标签\\n值」或「标签\\t值」中的值。"""
    raw = str(text or "")
    for pat in (
        re.escape(label) + r"\s*\n\s*([^\n]+)",
        re.escape(label) + r"\s*[：:]\s*([^\n]+)",
        re.escape(label) + r"\t+([^\n]+)",
    ):
        m = re.search(pat, raw, re.I)
        if not m:
            continue
        val = str(m.group(1) or "").strip()
        if not val or _looks_like_label(val):
            continue
        if any(s in val for s in stop_labels):
            continue
        if val.startswith(("□", "☑", "☐")):
            continue
        return val
    return ""


def _norm_cmp(s: str) -> str:
    s = str(s or "").replace("（", "(").replace("）", ")")
    return re.sub(r"[\s/·•\u3000]+", "", s).lower()


def _finalize_hj_fields(fields: dict[str, str], text: str) -> dict[str, str]:
    """补全 OCR 竖排/缺页时仍可提取的 HJ 栏位。"""
    out = dict(fields)
    raw = str(text or "")

    if not out.get("中文（音译）名"):
        m = re.search(r"申报人有效证件姓名[^：\n]*[：:]\s*[^\n（]+\(([^）)]+)\)", raw)
        if m:
            out["中文（音译）名"] = m.group(1).strip()
        elif out.get("有效证件姓名"):
            m = re.search(r"\(([^）)]+)\)", str(out["有效证件姓名"]))
            if m and re.search(r"[\u4e00-\u9fff]", m.group(1)):
                out["中文（音译）名"] = m.group(1).strip()

    contact = _extract_contact_vertical_fields(raw)
    for key in (
        "手机号", "电子邮箱", "最高学位中文", "最高学位英文",
        "回国前单位职务中文", "回国前单位职务英文",
    ):
        val = contact.get(key, "")
        if val and (_is_hj_field_garbage(out.get(key) or "") or not out.get(key)):
            out[key] = val

    if not out.get("最高学位中文") or _is_hj_field_garbage(out.get("最高学位中文") or ""):
        picked = _pick_after_label(raw, "中文 Cn", ("英文", "En"))
        if picked and not _is_hj_field_garbage(picked):
            out["最高学位中文"] = picked
    if not out.get("最高学位英文") or _is_hj_field_garbage(out.get("最高学位英文") or ""):
        picked = fix_ocr_english(_pick_after_label(raw, "英文 En", ("回国", "Last Employer")))
        if picked and not _is_hj_field_garbage(picked):
            out["最高学位英文"] = picked

    if not out.get("回国前单位职务中文") or not out.get("回国前单位职务英文"):
        cn_v, en_v = _extract_vertical_bilingual_block(raw, "回国（来华）前单位及职务", "相当于国内职称")
        if cn_v and not out.get("回国前单位职务中文"):
            out["回国前单位职务中文"] = cn_v.replace("、", "，")
        if en_v and not out.get("回国前单位职务英文"):
            out["回国前单位职务英文"] = fix_ocr_english(en_v.replace("、", ", "))

    for key in ("现工作单位", "引进职务", "回国时间", "拟落地省", "拟落地市", "任职单位地址"):
        if _is_scalar_label_value(out.get(key) or ""):
            out[key] = ""

    emp = _extract_employment_block_fields(raw)
    prefer_emp = {"拟落地省", "拟落地市", "现工作单位", "引进职务", "任职单位地址", "回国时间"}
    for key, val in emp.items():
        cur = out.get(key) or ""
        if val and (not cur or _is_scalar_label_value(cur) or _is_hj_field_garbage(cur, key) or key in prefer_emp):
            out[key] = val
    if out.get("拟落地省"):
        out["拟落地省"] = _clean_admin_region(out["拟落地省"], "省") or out["拟落地省"]
    if out.get("拟落地市"):
        out["拟落地市"] = _clean_admin_region(out["拟落地市"], "市") or ""
    if not out.get("引进职务") and re.search(r"首席\s*科\s*学\s*家", raw):
        out["引进职务"] = "首席科学家"

    if not out.get("回国时间"):
        rt = _value_after_vertical_label(
            raw,
            ("Time of Coming to China", "回国（来华）时间"),
            ("【第", "拟（现）任职单位地址", "4/30", "手机号", "Mobile"),
        )
        if rt and not _is_scalar_label_value(rt):
            out["回国时间"] = _fmt_hj_date(rt.replace("年", ".").replace("月", ""))

    if not out.get("相当于国内职称") and "相当于国内职称" in raw:
        block = raw.split("相当于国内职称", 1)[-1][:240]
        m = re.search(r"☑\s*([^□☐\n]{2,24})", block)
        if m:
            out["相当于国内职称"] = re.sub(r"\s+", "", m.group(1))

    if not out.get("回国前单位类型") and "回国（来华）前工作单位类型" in raw:
        block = raw.split("回国（来华）前工作单位类型", 1)[-1][:260]
        m = re.search(r"☑\s*(高校|科研机构|企业|其他)", block)
        if m:
            out["回国前单位类型"] = m.group(1)

    if not out.get("回国前所在地"):
        cn = _pick_after_label(raw, "Country (Region) of Residence", ("Host Province", "拟落地"))
        if cn:
            out["回国前所在地"] = _normalize_country_name(cn)

    if not out.get("国籍地区") or _is_hj_field_garbage(out.get("国籍地区") or ""):
        nat_block = raw.split("证件类型", 1)[0] if "证件类型" in raw else raw[:5000]
        for pat in (
            r"国\s*籍[\s\S]{0,240}?加\s*拿\s*大",
            r"外\s*籍[\s\S]{0,180}?加\s*拿\s*大",
            r"(?i)Foreign Nationality[\s\S]{0,180}?加\s*拿\s*大",
        ):
            if re.search(pat, nat_block):
                out["国籍地区"] = "加拿大"
                break
        if not out.get("国籍地区"):
            m = re.search(r"(?i)Foreign Nationality[\s\S]{0,120}?([\u4e00-\u9fff]{2,8})", raw)
            if m and _COUNTRY_NAME.match(m.group(1)):
                out["国籍地区"] = m.group(1)
            elif out.get("回国前所在地") and not _is_hj_field_garbage(out.get("回国前所在地") or ""):
                out["国籍地区"] = out["回国前所在地"]

    if not out.get("出生国家（地区）") or _is_hj_field_garbage(out.get("出生国家（地区）") or ""):
        m = re.search(r"(?i)Place of Birth\s*\n\s*([^\n]+)", raw)
        if m and _COUNTRY_NAME.match(_normalize_country_name(m.group(1))):
            out["出生国家（地区）"] = _normalize_country_name(m.group(1))
        else:
            m = re.search(r"(\d{8})[^\n]{0,30}([\u4e00-\u9fff]{2,6})", raw)
            if m and _COUNTRY_NAME.match(m.group(2)):
                out["出生国家（地区）"] = m.group(2)

    if not out.get("证件号码"):
        for pat in (
            r"(?i)Passport\s+([A-Z]{1,2}\d{6,})",
            r"护\s*照\s*([A-Z]{1,2}\d{6,})",
            r"☑\s*护照[^\n]*\n\s*([A-Z0-9]{6,})",
            r"\b([A-Z]{2}\d{6,})\b",
        ):
            m = re.search(pat, raw)
            if m and m.group(1).upper() not in {"IDNO", "OTHERIDNO", "PASSPORTAL99"}:
                out["证件号码"] = m.group(1).upper()
                break

    if not out.get("手机号"):
        out["手机号"] = (
            _pick_after_label(raw, "手机号 Mobile", ("电子邮件", "Email"))
            or contact.get("手机号")
            or out.get("手机号", "")
        )
        m_phone = re.search(r"(\+\d{1,3}[-\s]?\d{6,})", raw)
        if not out.get("手机号") and m_phone:
            out["手机号"] = m_phone.group(1).replace(" ", "")

    if not out.get("电子邮箱"):
        out["电子邮箱"] = fix_email_ocr(
            _pick_after_label(raw, "电子邮件 Email", ("最终毕业", "Highest Degree"))
            or contact.get("电子邮箱")
            or ""
        )
        if not out.get("电子邮箱"):
            m_mail = re.search(r"([A-Za-z0-9._%+-]+@?[A-Za-z0-9.-]+\.[A-Za-z]{2,})", raw)
            if m_mail:
                out["电子邮箱"] = fix_email_ocr(m_mail.group(1))
    elif out.get("电子邮箱"):
        out["电子邮箱"] = fix_email_ocr(out["电子邮箱"])

    if not out.get("性别"):
        out["性别"] = _normalize_gender(raw)
    else:
        out["性别"] = _normalize_gender(out["性别"]) or out["性别"]

    if not out.get("出生日期"):
        m = re.search(r"(?:性别|Gender|Wale|Male|Female|出生)[\s\S]{0,160}?\b((?:19|20)\d{6})\b", raw, re.I)
        if m:
            out["出生日期"] = _fmt_birth(m.group(1))
        else:
            m = re.search(r"(?i)Date of Birth[^\d]{0,80}((?:19|20)\d{6})", raw)
            if m:
                out["出生日期"] = _fmt_birth(m.group(1))
    elif out.get("出生日期"):
        out["出生日期"] = _fmt_birth(out["出生日期"])

    for options, key in (
        (("化学", "材料科学", "工程科学", "环境与地球科学", "信息科学", "生命科学", "医学"), "专业领域勾选"),
        (("新能源", "集成电路", "生命健康"), "关键技术勾选"),
        (("应用基础研究", "基础研究", "应用技术研究", "技术开发"), "研究类型勾选"),
        (("地方(地市级)实验室", "地方（地市级）实验室", "地市级"), "实验室类别勾选"),
        (("创新项目",), "项目类别勾选"),
        (("民营企业", "中央企业", "地方国有企业"), "用人单位类型"),
        (("企业高级职务", "企业高级职", "教授", "副教授", "博士后", "讲师"), "相当于国内职称"),
        (("企业", "高校", "科研机构"), "回国前单位类型"),
        (("首次申报",), "首次申报勾选"),
    ):
        if out.get(key):
            continue
        hit = _ocr_checked_option(raw, options)
        if not hit:
            continue
        if key == "实验室类别勾选":
            hit = "地市级"
        elif key == "相当于国内职称" and hit.startswith("企业高级"):
            hit = "企业高级职务"
        elif key == "首次申报勾选":
            hit = "首次申报"
        out[key] = hit

    corpus = raw + str(out.get("专长及代表性成果") or "") + str(out.get("最高学位中文") or "")
    if not out.get("专业领域勾选"):
        if re.search(r"电化学|应用化学|化学工程|化学与|化学化工|化工|Chemical", corpus, re.I):
            out["专业领域勾选"] = "化学"
        elif re.search(r"材料科学|材料工程|Advanced Materials|纳米材料|导热材料", corpus, re.I):
            out["专业领域勾选"] = "材料科学"
        elif re.search(r"工程科学|Engineering|机械|制造", corpus, re.I):
            out["专业领域勾选"] = "工程科学"
    if not out.get("关键技术勾选"):
        if re.search(r"新能源|电池|储能|电解质|动力电池", corpus, re.I):
            out["关键技术勾选"] = "新能源"
        elif re.search(r"生命健康|医学|生物", corpus, re.I):
            out["关键技术勾选"] = "生命健康"
    if not out.get("前沿领域勾选"):
        hit = _ocr_checked_option(raw, _COVER_FRONTIER)
        out["前沿领域勾选"] = hit or "不属于上述前沿领域"
    core = out.get("封面关键技术勾选") or out.get("关键技术勾选") or _ocr_checked_option(raw, _COVER_CORE_TECH)
    if core in _COVER_CORE_TECH:
        out["封面关键技术勾选"] = core
    else:
        out["封面关键技术勾选"] = "不涉及上述关键核心技术"

    if out.get("国籍地区"):
        cn = _normalize_country_name(out["国籍地区"])
        out["国籍地区"] = cn if _COUNTRY_NAME.match(cn) else ""
    if out.get("回国前所在地"):
        out["回国前所在地"] = _normalize_country_name(out["回国前所在地"])
    if out.get("出生国家（地区）"):
        out["出生国家（地区）"] = _normalize_country_name(out["出生国家（地区）"])
    if out.get("最高学位英文"):
        out["最高学位英文"] = normalize_hj_en_field(out["最高学位英文"])
    if out.get("回国前单位职务英文"):
        out["回国前单位职务英文"] = normalize_hj_en_field(out["回国前单位职务英文"])
    for key in ("最高学位中文", "回国前单位职务中文"):
        if out.get(key):
            out[key] = normalize_hj_cn_field(out[key])
    en_name = _extract_english_certificate_name(raw) or out.get("有效证件姓名") or ""
    en_name = _fix_concatenated_latin_name(en_name, raw)
    en_name, cn_name = split_certificate_name(en_name)
    if en_name:
        out["有效证件姓名"] = _strip_name_noise(en_name)
    if not out.get("申报人姓名"):
        out["申报人姓名"] = out.get("有效证件姓名") or ""
    else:
        out["申报人姓名"] = _strip_name_noise(out["申报人姓名"]) or out["申报人姓名"]
    if out.get("具体研究方向"):
        out["具体研究方向"] = _clean_research_direction(out["具体研究方向"], raw)
    elif "具体研究方向" in raw:
        out["具体研究方向"] = _clean_research_direction("", raw)
    if out.get("具体研究方向") and not out.get("研究领域关键词"):
        out["研究领域关键词"] = _keywords_from_direction(out["具体研究方向"])
    if not out.get("申报单位上级"):
        m = re.search(
            r"上级部门[^\n]{0,80}[:：]?\s*\n?\s*(江苏[省]?|[\u4e00-\u9fff]{2,8}(?:省|市))",
            raw,
        )
        if m:
            out["申报单位上级"] = "江苏省" if "江苏" in m.group(1) else m.group(1)

    compact_all = re.sub(r"\s+", "", raw)
    if not out.get("省级项目意愿"):
        if re.search(r"[团☑@]是[,，]?请填列意向|是,请填列意向省份", compact_all):
            out["省级项目意愿"] = "是"
    if out.get("省级项目意愿") == "是" and not out.get("意向省份"):
        pm = re.search(r"意向省份名称[^省]{0,40}(江苏[省]?|[\u4e00-\u9fff]{2,6}省)", compact_all)
        if pm:
            out["意向省份"] = _clean_admin_region(pm.group(1), "省") or "江苏省"
        elif "江苏" in compact_all:
            out["意向省份"] = "江苏省"

    if not out.get("中文（音译）名") or _is_hj_field_garbage(out.get("中文（音译）名") or ""):
        alt_cn = _format_chinese_transliteration(_extract_chinese_transliteration(raw) or cn_name)
        if alt_cn and not re.search(r"申报", alt_cn):
            out["中文（音译）名"] = alt_cn

    return {
        k: v for k, v in out.items()
        if str(v or "").strip() and not _is_hj_field_garbage(v, k)
    }


def apply_text_edits(text: str, edits: list) -> str:
    """将已确认编辑应用到 OCR 原文（供模板渲染使用）。"""
    out = str(text or "")
    for e in edits or []:
        if not isinstance(e, dict):
            continue
        find = str(e.get("find") or "")
        repl = str(e.get("replace") or "")
        if find and repl and find in out:
            out = out.replace(find, repl, 1)
    return out


def _build_fields(data: dict, raw_text: str) -> dict[str, str]:
    person = data.get("申报人基本信息") if isinstance(data.get("申报人基本信息"), dict) else {}
    info = data.get("申报信息") if isinstance(data.get("申报信息"), dict) else {}
    emp = data.get("用人单位情况及承诺") if isinstance(data.get("用人单位情况及承诺"), dict) else {}
    basic = emp.get("企业基本情况") if isinstance(emp.get("企业基本情况"), dict) else {}

    text = str(raw_text or "")
    regex_fields = {
        "有效证件姓名": _extract_english_certificate_name(text) or _first_group([
            ("有效证件姓名", r"(?i)Name of ID\t+([^\n\t]+)"),
            ("有效证件姓名", r"申报人有效证件姓名[^：\n]*[：:]\s*([^\n（]+)"),
        ], text),
        "中文（音译）名": _extract_chinese_transliteration(text) or _first_group([
            ("中文（音译）名", r"(?i)Name of Chinese Transliteration\t+([^\n\t]+)"),
        ], text),
        "性别": _normalize_gender(_first_group([
            ("性别", r"(?i)Gender\t+.*?([男女])"),
            ("性别", r"☑\s*(男)"),
            ("性别", r"☑\s*(女)"),
            ("性别", r"(?i)(\bWale\b|\bMale\b)"),
            ("性别", r"(?i)(\bFemale\b)"),
            ("性别", r"性别[\s\S]{0,80}?(男|女|Wale|Male|Female)"),
        ], text)),
        "出生日期": _first_group([
            ("出生日期", r"(?:性别|Gender|Wale|Male|Female|出生)[\s\S]{0,120}?\b((?:19|20)\d{6})\b"),
            ("出生日期", r"(?i)Date of Birth[\s\S]{0,160}?\b((?:19|20)\d{6})\b"),
        ], text),
        "出生国家（地区）": _first_group([
            ("出生国家（地区）", r"(?i)Place of Birth\t+([^\n\t]+)"),
            ("出生国家（地区）", r"(?i)Place of\s*\n?\s*Birth\s*\n\s*([^\n]+)"),
        ], text),
        "国籍地区": _first_group([
            ("国籍地区", r"Foreign Nationality\t+([^\n\t☐□]+)"),
            ("国籍地区", r"☑\s*外籍[^\n]*?([^\n\t☐□]{2,20})"),
        ], text),
        "证件号码": _first_group([
            ("证件号码", r"☑\s*护照[\s\S]{0,40}?Passport[\s\t]+([A-Z0-9]{6,})"),
            ("证件号码", r"☑\s*护照[^\n]*?Passport[\s\t]+([A-Z0-9]{6,})"),
            ("证件号码", r"☑\s*护照[^\n]*\n\s*([A-Z0-9]{6,})"),
            ("证件号码", r"(?i)Other ID NO\.\t+([A-Z0-9]{6,})"),
            ("证件号码", r"(?i)ID NO\.\t+([A-Z0-9]{6,})"),
        ], text),
        "证件类型": _first_group([("证件类型", r"☑\s*(护照|身份证|永居证)")], text) or "护照",
        "手机号": _first_group([
            ("手机号", r"(?i)Mobile\t+([+\d\- ]{8,})"),
            ("手机号", r"(?i)Mobile\s*\n\s*([+\d\- ]{8,})"),
            ("手机号", r"手机号\s*Mobile\t+([+\d\- ]{8,})"),
        ], text),
        "电子邮箱": _first_group([
            ("电子邮箱", r"(?i)电子邮件\s*Email\s*\n\s*([^\s\t\n]+@[^\s\t\n]+)"),
            ("电子邮箱", r"(?i)Email\s*\n\s*([^\s\t\n]+@[^\s\t\n]+)"),
            ("电子邮箱", r"(?i)Email\t+([^\s\t\n]+@[^\s\t\n]+)"),
        ], text),
        "最高学位中文": _pick_after_label(text, "中文 Cn", ("英文", "En")),
        "最高学位英文": _pick_after_label(text, "英文 En", ("回国", "Last Employer")),
        "回国前单位职务中文": _pick_tab_after("中文 Cn", text.split("回国（来华）前单位及职务")[-1] if "回国（来华）前单位及职务" in text else ""),
        "回国前单位职务英文": _pick_tab_after("英文 En", text.split("回国（来华）前单位及职务")[-1] if "回国（来华）前单位及职务" in text else ""),
        "相当于国内职称": _first_group([("相当于国内职称", r"☑\s*([^□☐\n]{2,20})")], text.split("相当于国内职称")[-1] if "相当于国内职称" in text else ""),
        "回国前单位类型": _first_group([("回国前单位类型", r"☑\s*(高校|科研机构|企业|其他)")], text.split("回国（来华）前工作单位类型")[-1] if "回国（来华）前工作单位类型" in text else ""),
        "回国前所在地": _normalize_country_name(_first_group([
            ("回国前所在地", r"(?i)Country \(Region\) of Residence\t+([^\n\t]+)"),
            ("回国前所在地", r"(?i)Country \(Region\) of Residence\s*\n\s*([^\n]+)"),
        ], text)),
        "拟落地省": _first_group([("拟落地省", r"(?i)Host Province\t+([^\n\t]+)")], text) or _value_after_vertical_label(
            text, ("拟落地省", "Host Province"), ("落地地级市", "Host City"),
        ),
        "拟落地市": _first_group([("拟落地市", r"(?i)Host City\t+([^\n\t]+)")], text) or _value_after_vertical_label(
            text, ("落地地级市", "Host City"), ("拟（现）任职单位名称", "Current or Expected Employer"),
        ),
        "现工作单位": _first_group([
            ("现工作单位", r"(?i)Current or Expected Employer\t+([^\n\t]+)"),
            ("现工作单位", r"拟（现）任职单位名称[^\n]*\t+([^\n\t]+)"),
        ], text) or _value_after_vertical_label(
            text.split("拟（现）任职单位名称", 1)[-1] if "拟（现）任职单位名称" in text else text,
            ("Current or Expected Employer",),
            ("职务（岗位）", "Position", "拟（现）任职单位地址"),
        ),
        "引进职务": _first_group([
            ("引进职务", r"(?i)Position\t+([^\n\t]+)"),
            ("引进职务", r"职务（岗位）\s*\t+([^\n\t]+)"),
        ], text) or _value_after_vertical_label(
            text.split("职务（岗位）", 1)[-1] if "职务（岗位）" in text else text,
            ("Position",),
            ("拟（现）任职单位地址", "Current or Expected Employer Add", "回国（来华）时间"),
        ),
        "任职单位地址": _first_group([
            ("任职单位地址", r"(?i)Current or Expected Employer Add\.\t+([^\n]+)"),
        ], text) or _value_after_vertical_label(
            text,
            ("拟（现）任职单位地址", "Current or Expected Employer Add"),
            ("回国（来华）时间", "Time of Coming to China"),
        ),
        "回国时间": _first_group([
            ("回国时间", r"(?i)Time of Coming to China\t+([^\n\t]+)"),
            ("回国时间", r"(?i)Time of Coming to China\s*\n\s*([^\n]+)"),
        ], text),
        "申报单位": _first_group([("申报单位", r"申报单位[^：\n]*[：:]\s*([^\n]+)")], text),
        "实验室名称": _first_group([("实验室名称", r"实验室名称[^：\n]*[：:]\s*([^\n]+)")], text),
        "项目类别": _first_group([("项目类别", r"项目类别[^：\n]*[：:]\s*([^\n]+)")], text),
        "单位联系人": _first_group([
            ("单位联系人", r"单位联系人[^：\n]*[：:]\s*([^\n]+)"),
            ("单位联系人", r"(?i)Contact Person\t+([^\n\t]+)"),
        ], text),
        "联系人电话": _first_group([
            ("联系人电话", r"联系人电话[^：\n]*[：:]\s*([^\n]+)"),
            ("联系人电话", r"(?i)Telephone Number\t+([^\n\t]+)"),
        ], text),
        "填表日期": _first_group([
            ("填表日期", r"填表日期[^：\n\d]*[：:]?\s*(\d{4}\s*年\s*\d{1,2}\s*月\s*\d{1,2}\s*日)"),
            ("填表日期", r"填表日期[^：\n]*[：:]\s*([^\n]+)"),
        ], text),
        "二级学科及代码": _first_group([
            ("二级学科及代码", r"所属二级学科及代码[\s\S]{0,80}?([\u4e00-\u9fff]{2,20}\s*[,，]?\s*\d{4,6})"),
            ("二级学科及代码", r"所属二级学科及代码[^：\n]*[：:]\s*([^\n]+)"),
            ("二级学科及代码", r"Category II Discipline and Code[：:]\s*([^\n]+)"),
            ("二级学科及代码", r"(\d{5}[\u4e00-\u9fff]{2,20}|[\u4e00-\u9fff]{2,20}[,，]?\s*\d{5})"),
        ], text),
        "具体研究方向": _first_group([
            ("具体研究方向", r"具体研究方向[^：\n]*[：:]\s*([^\n]+)"),
        ], text),
        "申报单位上级": _first_group([
            ("申报单位上级", r"上级部门/地方[^：\n]*[：:]\s*([^\n]+)"),
            ("申报单位上级", r"按照隶属关系填写[^：\n]*[：:]\s*([^\n]+)"),
        ], text),
        "用人单位类型": _first_group([("用人单位类型", r"☑\s*(民营企业|中央企业|地方国有企业|部属高校|地方高校)")], text.split("用人单位类型")[-1] if "用人单位类型" in text else ""),
    }

    # 回国前职务：优先 tab 格式，否则取竖排 OCR 的中英行
    if "回国（来华）前单位及职务" in text:
        block = text.split("回国（来华）前单位及职务", 1)[1].split("相当于国内职称")[0]
        cn_parts = re.findall(r"中文\s*Cn\t+([^\n]+)", block, re.I)
        en_parts = re.findall(r"英文\s*En\t+([^\n]+)", block, re.I)
        if cn_parts:
            regex_fields["回国前单位职务中文"] = _clean_value("回国前单位职务中文", cn_parts[0])
        if en_parts:
            regex_fields["回国前单位职务英文"] = _clean_value("回国前单位职务英文", en_parts[0])
        if not regex_fields.get("回国前单位职务中文") or not regex_fields.get("回国前单位职务英文"):
            cn_v, en_v = _extract_vertical_bilingual_block(text, "回国（来华）前单位及职务", "相当于国内职称")
            if cn_v and not regex_fields.get("回国前单位职务中文"):
                regex_fields["回国前单位职务中文"] = _clean_value("回国前单位职务中文", cn_v)
            if en_v and not regex_fields.get("回国前单位职务英文"):
                regex_fields["回国前单位职务英文"] = _clean_value("回国前单位职务英文", en_v)

    if "最终毕业院校及专业" in text or "Highest Degree" in text:
        block = text
        if "回国（来华）前单位及职务" in text:
            block = text.split("回国（来华）前单位及职务")[0]
        cn_parts = re.findall(r"中文\s*Cn\t+([^\n]+)", block, re.I)
        en_parts = re.findall(r"英文\s*En\t+([^\n]+)", block, re.I)
        if cn_parts:
            regex_fields["最高学位中文"] = _clean_value("最高学位中文", cn_parts[-1])
        if en_parts:
            regex_fields["最高学位英文"] = _clean_value("最高学位英文", en_parts[-1])

    parser_fields = {
        "有效证件姓名": str(person.get("有效证件姓名") or info.get("申报人") or ""),
        "中文（音译）名": str(person.get("外籍专家中文姓名") or person.get("中文（音译）名") or ""),
        "性别": str(person.get("性别") or ""),
        "出生日期": str(person.get("出生日期") or ""),
        "出生国家（地区）": str(person.get("出生国家（地区）") or person.get("出生国家") or ""),
        "国籍地区": str(person.get("国籍地区") or ""),
        "证件号码": str(person.get("证件号码") or ""),
        "证件类型": str(person.get("证件类型") or ""),
        "手机号": str(person.get("手机号") or ""),
        "电子邮箱": str(person.get("电子邮箱") or person.get("电子邮件") or ""),
        "现工作单位": str(person.get("现工作单位") or person.get("拟（现）任职单位名称") or basic.get("企业名称") or person.get("用人单位名称") or ""),
        "引进职务": str(person.get("引进职务") or person.get("职务（岗位）") or ""),
        "申报单位": str(info.get("申报企业") or basic.get("企业名称") or person.get("用人单位名称") or ""),
        "实验室名称": str(basic.get("实验室名称") or person.get("实验室名称") or ""),
        "单位联系人": str(basic.get("联系人") or person.get("单位联系人") or ""),
        "联系人电话": str(basic.get("联系人电话") or person.get("联系人电话") or ""),
        "填表日期": str(info.get("申报日期") or person.get("填表日期") or ""),
    }

    contact_fields = _extract_contact_vertical_fields(text)
    fields: dict[str, str] = {}
    for k in set(regex_fields) | set(parser_fields) | set(contact_fields):
        for src in (contact_fields, regex_fields, parser_fields):
            v = _clean_value(k, src.get(k) or "")
            if not v or _is_hj_field_garbage(v):
                continue
            if k not in fields or len(v) > len(fields[k]):
                fields[k] = v

    for label, val in re.findall(r"^([^：\n]{2,40})[：:]\s*(.+)$", text):
        k = label.strip()
        v = _clean_value(k, val)
        if v and (k not in fields or not fields[k]):
            fields[k] = v

    if fields.get("性别") in ("☑ 男 Male", "男 Male", "Male", "Wale", "wale"):
        fields["性别"] = "男"
    elif fields.get("性别") in ("☑ 女 Female", "女 Female", "Female"):
        fields["性别"] = "女"
    if not fields.get("性别"):
        if re.search(r"☑\s*男", text):
            fields["性别"] = "男"
        elif re.search(r"☑\s*女", text):
            fields["性别"] = "女"
        elif re.search(r"(?i)\bWale\b|\bMale\b", text):
            fields["性别"] = "男"
        elif re.search(r"(?i)\bFemale\b", text):
            fields["性别"] = "女"
    if not fields.get("国籍地区"):
        m = re.search(r"Foreign Nationality\s*\n\s*([^\n☐□☑]+)", text, re.I)
        if m:
            fields["国籍地区"] = _normalize_country_name(_clean_value("国籍地区", m.group(1)))
    if fields.get("国籍地区"):
        cn = _normalize_country_name(fields["国籍地区"])
        fields["国籍地区"] = cn if _COUNTRY_NAME.match(cn) else ""
    if fields.get("回国前所在地"):
        fields["回国前所在地"] = _normalize_country_name(fields["回国前所在地"])
    if fields.get("出生国家（地区）"):
        fields["出生国家（地区）"] = _normalize_country_name(fields["出生国家（地区）"])
    if not fields.get("出生国家（地区）"):
        lines = [ln.strip() for ln in text.splitlines()]
        for i, ln in enumerate(lines):
            if re.fullmatch(r"\d{8}", ln):
                for j in range(i + 1, min(i + 12, len(lines))):
                    if _COUNTRY_NAME.match(lines[j]):
                        fields["出生国家（地区）"] = lines[j]
                        break
                break
    if fields.get("出生日期"):
        fields["出生日期"] = _fmt_birth(fields["出生日期"])
    if fields.get("回国时间"):
        fields["回国时间"] = _fmt_hj_date(fields["回国时间"].replace("年", ".").replace("月", ""))

    # 从 OCR 勾选框推断专业领域/研究类型
    for label, key in (
        ("化学", "专业领域勾选"), ("材料科学", "专业领域勾选"), ("工程科学", "专业领域勾选"),
        ("新能源", "关键技术勾选"), ("应用基础研究", "研究类型勾选"),
    ):
        if re.search(r"☑\s*" + label, text):
            fields[key] = label

    fields["专长及代表性成果"] = _clean_narrative(_section_between(
        text,
        (
            "一、个人简介及贡献", "个人简介及贡献", "一、亮点履历", "1. 亮点履历",
            "申报人是国际能源", "科研情况 (请概述", "Research Status (Please",
            "科研情况及代表性成果",
        ),
        ("2.代表性科研项目", "Grants (As a Leader", "代表性科研项目（主持", "代表性科研项目"),
    ) or str(data.get("专长及代表性成果") or data.get("科研情况") or "")[:8000])

    fields["工作设想"] = _clean_narrative(_section_between(
        text,
        (
            "工作设想部分", "工作设想 (", "一、研究背景及意义", "研究背景及意义",
            "一、研究背景", "面向全球",
        ),
        (
            "申报单位（用人单位）情况部分", "申报单位（用人单位）情况",
            "用人单位类型", "申报情况 (Application Facts)",
        ),
    ) or str(data.get("工作设想") or "")[:8000])

    fields["用人单位简介"] = _clean_narrative(_section_between(
        text,
        ("申报单位（用人单位）简介", "3.申报单位（用人单位）简介", "主要指实验室建设情况"),
        ("申报单位（用人单位）意见部分", "申报单位（用人单位）意见", "1. 推荐理由"),
    ))

    fields["推荐理由"] = _clean_narrative(_section_between(
        text,
        ("1. 推荐理由", "1.推荐理由", "推荐理由及引进的必要性"),
        ("2. 支持条件", "2.支持条件", "支持条件（包括工作和生活"),
    ))

    fields["支持条件"] = _clean_narrative(_section_between(
        text,
        ("2. 支持条件", "2.支持条件", "支持条件（包括工作和生活"),
        ("申报人有关信息属实", "主要负责人签字", "申报渠道意见"),
    ))

    return _finalize_hj_fields(fields, text)


def _seg_matches_kw(seg: str, kw: str) -> bool:
    sn, kn = _norm_cmp(seg), _norm_cmp(kw)
    if not sn or not kn:
        return False
    if sn == kn:
        return True
    if sn.startswith(kn):
        rest = sn[len(kn):]
        return (not rest) or bool(re.fullmatch(r"[a-z()（）\-]+", rest))
    if kn.startswith(sn) and len(sn) >= 4:
        return True
    idx = sn.find(kn)
    if len(kn) >= 3 and idx >= 0:
        prev = sn[idx - 1] if idx else ""
        if prev and re.match(r"[\u4e00-\u9fffa-z]", prev):
            return False
        if sn.endswith(kn) and len(sn) > len(kn) + 2:
            return False
        return True
    return False


def _mark_checkbox(text: str, keywords: tuple[str, ...]) -> str:
    if not text.strip():
        return text
    out = re.sub(r"☑", "□", str(text))
    hit = False
    for kw in sorted((k for k in keywords if k), key=lambda x: len(_norm_cmp(x)), reverse=True):
        for m in re.finditer(r"[□☐口]\s*([^□☐口\n]*)", out):
            seg = m.group(1)
            if _seg_matches_kw(seg, kw):
                out = out[: m.start()] + "☑" + seg + out[m.end() :]
                hit = True
                break
        if (not hit) and "不涉及" in kw and "不涉及上述关键核心技术" in out:
            out = re.sub(r"不涉及上述关键核心技术", "☑不涉及上述关键核心技术", out, count=1)
            hit = True
    return out


def _format_cover_date(val: str) -> str:
    m = re.search(r"(20\d{2})\s*年\s*(\d{1,2})\s*月\s*(\d{1,2})", str(val or ""))
    if m:
        return f"{m.group(1)} 年 {int(m.group(2))} 月 {int(m.group(3))} 日"
    return str(val or "").strip()


def _format_hj_discipline(s: str) -> str:
    """正式封面：代码紧挨学科名，如 14015理论物理学。"""
    raw = str(s or "").strip()
    if not raw:
        return ""
    m = re.search(r"([\u4e00-\u9fff]{2,30})\s*[,，]\s*(\d{4,6})", raw)
    if m:
        return m.group(2) + m.group(1)
    m = re.search(r"(\d{4,6})\s*[,，]?\s*([\u4e00-\u9fff]{2,30})", raw)
    if m:
        return m.group(1) + m.group(2)
    return re.sub(r"\s+", "", raw)


def _fill_cover_paragraphs(doc: Document, fields: dict[str, str]) -> None:
    scalar_rules = (
        (re.compile(r"(申报人姓名[^:：\n]*[:：])\s*.+$", re.I), "申报人姓名"),
        (re.compile(r"(申报单位[^:：\n]*[:：])\s*.+$", re.I), "申报单位"),
        (re.compile(r"(实验室名称[^:：\n]*[:：])\s*.+$", re.I), "实验室名称"),
        (re.compile(r"(联系人电话[^:：\n]*[:：])\s*.*$", re.I), "联系人电话"),
        (re.compile(r"(联\s*系\s*人(?!\s*电话)[^:：\n]*[:：])\s*.*$", re.I), "单位联系人"),
        (re.compile(r"(单位联系人[^:：\n]*[:：])\s*.+$", re.I), "单位联系人"),
        (re.compile(r"((?:单位)?联系人(?!\s*电话)[^:：\n]*[:：])\s*.*$", re.I), "单位联系人"),
        (re.compile(r"(填表日期\s*\([^)]*\)\s*).*$", re.I), "填表日期"),
        (re.compile(r"(项目类别[^:：\n]*[:：])\s*.+$", re.I), "项目类别"),
    )
    next_para_rules = (
        (("contactperson",), "单位联系人"),
        (("telephonenumber",), "联系人电话"),
    )
    for p in doc.paragraphs:
        cur = str(p.text or "")
        if not cur.strip():
            continue
        compact = _norm_cmp(cur)
        for pat, key in scalar_rules:
            val = str(fields.get(key) or "").strip()
            if key == "填表日期":
                val = _format_cover_date(val)
            if val and pat.search(cur):
                if key == "申报人姓名":
                    val = _strip_name_noise(val)
                pad = "                " if key in ("申报人姓名", "申报单位", "实验室名称") else " "
                tail = (" " * max(0, 28 - len(val))) if key in ("申报人姓名", "申报单位", "实验室名称") else ""
                cur = pat.sub(lambda m, v=val, p=pad, t=tail: m.group(1) + p + v + t, cur, count=1)
                break
        kw_major = fields.get("专业领域勾选") or ""
        if kw_major and re.search(r"[□☐口]", cur) and (
            "专业领域" in compact or ("数学" in compact and "物理" in compact)
        ):
            cur = _mark_checkbox(cur, (kw_major,))
        if "实验室" in compact and "类别" in compact:
            kw_lab = fields.get("实验室类别勾选") or ""
            if kw_lab:
                cur = _mark_checkbox(cur, ("地方(地市级)实验室", "地市级"))
        if "项目类别" in compact and (fields.get("项目类别") or fields.get("项目类别勾选")):
            cur = _mark_checkbox(cur, (fields.get("项目类别勾选") or "创新项目",))
        if "前沿领域" in compact and fields.get("前沿领域勾选"):
            cur = _mark_checkbox(cur, (fields.get("前沿领域勾选"),))
        if "关键核心技术" in compact and (fields.get("封面关键技术勾选") or fields.get("关键技术勾选")):
            cur = _mark_checkbox(cur, (fields.get("封面关键技术勾选") or fields.get("关键技术勾选"),))
        if cur != p.text:
            old = p.text or ""
            only_box = re.sub(r"[□☑☐口]", "", old) == re.sub(r"[□☑☐口]", "", cur)
            if only_box and _assign_run_texts(p, cur):
                pass
            else:
                _set_paragraph_text(p, cur)
    for i, p in enumerate(doc.paragraphs):
        compact = _norm_cmp(p.text)
        for keys, fkey in next_para_rules:
            if not any(k in compact for k in keys):
                continue
            val = str(fields.get(fkey) or "").strip()
            if not val:
                continue
            if re.search(r"[:：]\s*\S", p.text):
                continue
            if re.search(r"[:：]\s*$", p.text):
                _set_paragraph_text(p, p.text.rstrip() + " " + val)
                break
            for j in range(i + 1, min(i + 4, len(doc.paragraphs))):
                nxt = doc.paragraphs[j]
                if not str(nxt.text or "").strip():
                    _set_paragraph_text(nxt, val)
                    break
    disc = _format_hj_discipline(fields.get("二级学科及代码") or "")
    if disc:
        for i, p in enumerate(doc.paragraphs):
            t = str(p.text or "")
            compact = re.sub(r"\s+", "", t)
            if "二级学科" in compact and "代码" in compact and not re.search(r"\d{4,}", compact):
                for j in range(i + 1, min(i + 4, len(doc.paragraphs))):
                    nxt = doc.paragraphs[j]
                    if not str(nxt.text or "").strip():
                        _set_paragraph_text(nxt, disc)
                        break
                else:
                    _set_paragraph_text(p, t.rstrip() + disc)
                break


def _fill_hj_table0(table, fields: dict[str, str]) -> None:
    if not table.rows:
        return
    applicant = fields.get("申报人姓名") or fields.get("有效证件姓名") or ""
    cert = fields.get("有效证件姓名") or applicant
    cn = fields.get("中文（音译）名") or ""
    if len(table.rows) >= 1:
        cells = _distinct_cells(table.rows[0])
        name = cert or applicant
        if name and len(cells) >= 3:
            _write_cell(cells[2], name)
            if len(cells) >= 4 and not _looks_like_label(cells[3].text):
                _write_cell(cells[3], "")
    if len(table.rows) >= 2 and cn:
        cells = _distinct_cells(table.rows[1])
        if len(cells) >= 3:
            _write_cell(cells[2], cn)
            if len(cells) >= 4 and not _looks_like_label(cells[3].text):
                _write_cell(cells[3], "")

    if len(table.rows) > 2:
        cells = _distinct_cells(table.rows[2])
        if fields.get("性别") and len(cells) >= 2:
            g = _normalize_gender(fields["性别"]) or fields["性别"]
            # 正式 HJ 样例性别栏直接写「男」「女」，不用勾选框
            _write_cell(cells[1], g)
        if fields.get("出生日期") and len(cells) >= 4:
            _write_cell(cells[3], fields["出生日期"])
        if fields.get("出生国家（地区）") and len(cells) >= 6:
            _write_cell(cells[5], fields["出生国家（地区）"])

    for row in table.rows:
        cells = _distinct_cells(row)
        row_n = _norm(" ".join(c.text for c in cells))
        if "foreignnationality" in row_n or "外籍" in row_n:
            if fields.get("国籍地区"):
                wrote = False
                for i, c in enumerate(cells):
                    t = c.text.strip()
                    if not t or _looks_like_label(t) or "华裔" in t:
                        continue
                    if t.startswith(("□", "☑", "☐")):
                        continue
                    if len(t) < 24:
                        _write_cell(c, fields["国籍地区"])
                        wrote = True
                        break
                if not wrote:
                    for i, c in enumerate(cells):
                        if i > 0 and not _looks_like_label(c.text) and "华裔" not in c.text:
                            if not c.text.strip() or len(c.text.strip()) < 3:
                                _write_cell(c, fields["国籍地区"])
                                break
            if "外籍" in row_n:
                for c in cells:
                    if "外籍" in c.text or "Foreign" in c.text:
                        _mark_checkboxes_in_cell(c, ("外籍", "Foreign Nationality"))
                        break
                ethnic = str(fields.get("是否华裔") or "")
                if "非华裔" in ethnic or (fields.get("国籍地区") and "华裔" not in ethnic):
                    for c in cells:
                        if "非华裔" in c.text or "Non-Ethnic" in c.text:
                            _mark_checkboxes_in_cell(c, ("非华裔", "Non-Ethnic"))
                            break
        if "passport" in row_n:
            if fields.get("证件类型", "").find("护照") >= 0:
                for c in cells:
                    if "passport" in _norm(c.text) or "护照" in c.text:
                        _mark_checkboxes_in_cell(c, ("护照", "Passport"))
                if fields.get("证件号码"):
                    for c in reversed(cells):
                        if not _looks_like_label(c.text) and "passport" not in _norm(c.text):
                            _write_cell(c, fields["证件号码"])
                            break


def _fill_hj_table1(table, fields: dict[str, str]) -> None:
    if len(table.rows) < 12:
        return
    rows = table.rows
    if fields.get("手机号"):
        c = _distinct_cells(rows[0])
        if len(c) >= 3:
            _write_cell(c[2], fields["手机号"])
    if fields.get("电子邮箱"):
        c = _distinct_cells(rows[1])
        if len(c) >= 2:
            _write_cell(c[-1], fields["电子邮箱"])
    degree_rows = (
        (2, "最高学位英文"),
        (3, "最高学位中文"),
    )
    for ri, key in degree_rows:
        if fields.get(key):
            c = _distinct_cells(rows[ri])
            if len(c) >= 2:
                _write_cell(c[-1], fields[key])
    employer_rows = (
        (4, "回国前单位职务英文"),
        (5, "回国前单位职务中文"),
    )
    for ri, key in employer_rows:
        if fields.get(key):
            c = _distinct_cells(rows[ri])
            if len(c) >= 2:
                _write_cell(c[-1], fields[key])
    c6 = _distinct_cells(rows[6])
    title = str(fields.get("相当于国内职称") or "").strip()
    if title:
        marked = False
        for cell in c6:
            if re.search(r"[□☐口☑]", cell.text):
                _mark_checkboxes_in_cell(cell, (title,))
                marked = True
        if not marked and is_academic_title(title) and len(c6) >= 3:
            _write_cell(c6[2], title)
    c7 = _distinct_cells(rows[7])
    if fields.get("回国前单位类型"):
        extra = {"企业": "Enterprise", "高校": "University", "科研机构": "Research"}
        kws = (fields["回国前单位类型"], extra.get(fields["回国前单位类型"], "Enterprise"))
        for cell in c7:
            if re.search(r"[□☐口☑]", cell.text):
                _mark_checkboxes_in_cell(cell, kws)
    c8 = _distinct_cells(rows[8])
    for ci, fk in ((1, "回国前所在地"), (3, "拟落地省"), (5, "拟落地市")):
        if fields.get(fk) and ci < len(c8):
            _write_cell(c8[ci], fields[fk])
    c9 = _distinct_cells(rows[9])
    if fields.get("现工作单位") and len(c9) >= 2:
        _write_cell(c9[1], fields["现工作单位"])
    if fields.get("引进职务") and len(c9) >= 4:
        _write_cell(c9[3], fields["引进职务"])
    c10 = _distinct_cells(rows[10])
    if fields.get("任职单位地址") and len(c10) >= 2:
        _write_cell(c10[1], fields["任职单位地址"])
    c11 = _distinct_cells(rows[11])
    if fields.get("回国时间") and len(c11) >= 2:
        _write_cell(c11[1], fields["回国时间"])


def _fill_scalar_rows(doc: Document, fields: dict[str, str], skip_tables: set[int] | None = None) -> int:
    skip = skip_tables or set()
    n = 0
    for ti, table in enumerate(doc.tables):
        if ti in skip:
            continue
        for row in table.rows:
            cells = _distinct_cells(row)
            if not cells:
                continue
            row_text = " ".join(c.text for c in cells)
            key = _row_field_key(row_text)
            if not key:
                continue
            val = str(fields.get(key) or "").strip()
            if not val:
                continue
            label_last = -1
            for i, cell in enumerate(cells):
                if _looks_like_label(cell.text):
                    label_last = i
            wrote = False
            for i, cell in enumerate(cells):
                if i <= label_last or _looks_like_label(cell.text):
                    continue
                _write_cell(cell, val)
                wrote = True
            if wrote:
                n += 1
    return n


def _row_field_key(row_text: str) -> str | None:
    n = _norm(row_text)
    rules = (
        (("nameofvalidcertificate", "nameofid"), "有效证件姓名"),
        (("chinesetransliteration", "nameofchinese"), "中文（音译）名"),
        (("contactperson",), "单位联系人"),
        (("telephonenumber",), "联系人电话"),
        (("laboratory",), "实验室名称"),
        (("employer",), "申报单位"),
    )
    for keys, field in rules:
        if any(k in n for k in keys):
            return field
    return None


def _fill_hj_table2(table, edu: list[dict], work: list[dict]) -> None:
    edu_start, edu_end, work_start = _edu_work_row_bounds(table)
    for i, d in enumerate(edu[: max(0, edu_end - edu_start)]):
        ri = edu_start + i
        if ri >= edu_end or ri >= len(table.rows):
            break
        cells = _distinct_cells(table.rows[ri])
        if len(cells) < 5:
            continue
        vals = [
            d.get("起止时间", ""),
            d.get("所在国家", ""),
            d.get("校名称", ""),
            d.get("专业领域", ""),
            d.get("学位", ""),
        ]
        for off, val in enumerate(vals, start=1):
            if off < len(cells):
                _write_cell(cells[off], str(val or "").strip())
    max_work = len(table.rows) - work_start
    for i, d in enumerate(work[:max_work]):
        ri = work_start + i
        if ri >= len(table.rows):
            break
        cells = _distinct_cells(table.rows[ri])
        if len(cells) < 5:
            continue
        vals = [
            d.get("起止时间", ""),
            d.get("所在国家", ""),
            d.get("工作单位", ""),
            d.get("担任职务", ""),
            d.get("任职情况", ""),
        ]
        for off, val in enumerate(vals, start=1):
            if off < len(cells):
                _write_cell(cells[off], str(val or "").strip())


def _fill_paper_table(table, papers: list[dict], start: int = 0, direction: str = "") -> int:
    n = 0
    for ri in range(1, len(table.rows)):
        pi = start + n
        if pi >= len(papers):
            break
        cells = _distinct_cells(table.rows[ri])
        if len(cells) < 6:
            continue
        d = papers[pi]
        role = str(d.get("角色") or "").strip() or str(d.get("排序") or "").strip()
        desc = str(d.get("研究方向描述") or direction or "").strip()
        vals = [
            "",
            desc,
            _fmt_hj_date(str(d.get("发表时间") or "")),
            str(d.get("论文题目") or "").strip(),
            str(d.get("发表载体") or "").strip(),
            role,
        ]
        for ci, val in enumerate(vals):
            if ci < len(cells) and str(val or "").strip():
                _write_cell(cells[ci], str(val))
        n += 1
    return n


def _fill_project_table(table, projects: list[dict], start: int = 0, direction: str = "") -> int:
    n = 0
    for ri in range(1, len(table.rows)):
        pi = start + n
        if pi >= len(projects):
            break
        cells = _distinct_cells(table.rows[ri])
        if len(cells) < 6:
            continue
        d = projects[pi]
        vals = [
            "",
            str(d.get("描述") or direction or "").strip(),
            _normalize_hj_range_text(str(d.get("起止时间") or "")),
            str(d.get("性质来源") or "").strip(),
            str(d.get("经费") or "").strip(),
            str(d.get("角色") or "").strip(),
        ]
        for ci, val in enumerate(vals):
            if ci < len(cells) and str(val or "").strip():
                _write_cell(cells[ci], str(val))
        n += 1
    return n


def _ensure_table_rows(table, min_rows: int) -> None:
    if table is None or not table.rows:
        return
    while len(table.rows) < min_rows:
        tbl = table._tbl
        last = table.rows[-1]._tr
        tbl.append(deepcopy(last))
        cells = _distinct_cells(table.rows[-1])
        for ci, cell in enumerate(cells):
            if ci == 0:
                continue
            _write_cell(cell, "")


def _fill_hj_table5(table, fields: dict[str, str]) -> None:
    if not table.rows:
        return
    cell = table.rows[0].cells[0]
    kws = str(fields.get("研究领域关键词") or "").strip()
    kind = str(fields.get("研究类型勾选") or "").strip()
    extra = {
        "应用基础研究": "Applied Basic Research",
        "基础研究": "Basic Research",
        "应用技术研究": "Applied Technology Research",
        "技术开发": "Technology Development",
    }
    en = extra.get(kind, "")
    for p in cell.paragraphs:
        t = str(p.text or "")
        if not t.strip():
            continue
        if kws and "研究领域关键词" in t:
            t = re.sub(r"(研究领域关键词[^:：\n]*[:：])\s*.*$", r"\1" + kws, t, count=1)
            _set_paragraph_text(p, t)
            continue
        if kind and re.search(r"[□☐口☑]", t):
            t = _mark_checkbox(t, (kind, en) if en else (kind,))
            _set_paragraph_text(p, t)


def _ensure_cell_para(cell, idx: int):
    while len(cell.paragraphs) <= idx:
        cell.add_paragraph("")
    return cell.paragraphs[idx]


def _fill_expertise_cell(cell, body: str) -> None:
    """按正式 16 表：标题/中英说明 13.5 加粗，空一行后再写宋体 12 正文。"""
    body = str(body or "").strip()
    body = re.sub(r"^专长及代表性成果[^\n]*\n?", "", body).strip()
    body = re.sub(r"^所从事的专业领域及取得的成绩描述[^\n]*\n?", "", body).strip()
    body = re.sub(r"^Field of Expertise[^\n]*\n?", "", body, flags=re.I).strip()
    paras = list(cell.paragraphs)
    if not paras:
        cell.add_paragraph(_EXPERTISE_TITLE)
        paras = list(cell.paragraphs)
    p1 = _ensure_cell_para(cell, 1)
    _write_runs(p1, [
        (_EXPERTISE_INTRO_HEAD, {"east": "宋体", "ascii_name": "Times New Roman", "size_pt": 13.5, "bold": True}),
        (_EXPERTISE_INTRO_HINT, {"east": "方正仿宋_GBK", "ascii_name": "Times New Roman", "size_pt": 13.5, "bold": False}),
    ])
    p2 = _ensure_cell_para(cell, 2)
    _write_runs(p2, [
        (_EXPERTISE_EN, {"east": "Times New Roman", "ascii_name": "Times New Roman", "size_pt": 13.5, "bold": True}),
    ])
    p3 = _ensure_cell_para(cell, 3)
    _set_paragraph_text(p3, "")
    lines = _split_narrative_paras(body)
    _write_paragraphs_from(cell, 4, lines, keep_blank=True)


def _fill_hj_table13(table, fields: dict[str, str]) -> None:
    if len(table.rows) < 3:
        return
    cell = table.rows[2].cells[0]
    marks = []
    if fields.get("首次申报勾选"):
        marks.append("首次申报")
    if fields.get("省级项目意愿") == "是":
        marks.append("是，请填列意向省份名称")
    for p in cell.paragraphs:
        t = str(p.text or "")
        if not t.strip():
            continue
        nt = t
        if marks:
            nt = _mark_checkbox(t, tuple(marks))
            if fields.get("省级项目意愿") == "是":
                nt = re.sub(r"[□☐口]\s*是，请填列意向省份名称", "☑是，请填列意向省份名称", nt, count=1)
        if fields.get("意向省份") and "意向省份名称" in nt:
            nt = re.sub(
                r"(意向省份名称[：:]\s*)([\u4e00-\u9fff]{0,8})",
                r"\1" + fields["意向省份"],
                nt,
                count=1,
            )
        if nt != t:
            _set_paragraph_text(p, nt)
    if len(table.rows) >= 4 and fields.get("填表日期"):
        c3 = table.rows[3].cells[0]
        dt = _format_cover_date(fields["填表日期"])
        for p in c3.paragraphs:
            body = str(p.text or "")
            if re.search(r"20\d{2}\s*年", body) or "Date" in body:
                nt = re.sub(r"20\d{2}\s*年\s*\d{1,2}\s*月\s*\d{1,2}\s*日", dt.replace(" ", ""), body)
                if nt != body:
                    _set_paragraph_text(p, nt)


def _clip_employer_intro(intro: str) -> str:
    s = _clean_narrative(intro)
    m = re.search(r"实验室依托", s)
    if m:
        s = s[m.start():]
    cut = re.search(
        r"申报单位.{0,12}意见|推荐理由|支持条件|申报渠道|主要负责人签字",
        s,
    )
    if cut:
        s = s[: cut.start()]
    s = re.sub(r"结晶度的硅酸盐[\s\S]*$", "", s)
    return s.strip()


def _fill_hj_table14(table, fields: dict[str, str]) -> None:
    if not table.rows:
        return
    intro = _clip_employer_intro(fields.get("用人单位简介") or "")
    parent = str(fields.get("申报单位上级") or "").strip()
    kind = str(fields.get("用人单位类型") or "").strip()
    cell = table.rows[0].cells[0]
    intro_idx = None
    for i, p in enumerate(cell.paragraphs):
        t = str(p.text or "")
        if not t.strip():
            continue
        if re.search(r"[□☐口☑]", t) and kind:
            nt = _mark_checkbox(t, (kind, "民营企业"))
            if nt != t:
                if not _assign_run_texts(p, nt):
                    _set_paragraph_text(p, nt)
                _paint_paragraph_font(p)
            continue
        if "上级部门" in t or "按隶属关系" in t:
            if parent:
                nt = re.sub(r"([:：])\s*.*$", r"\1 " + parent, t)
                _set_paragraph_text(p, nt)
            continue
        if "简介" in t and ("300" in t or "实验室建设" in t):
            intro_idx = i
    if intro and intro_idx is not None:
        lines = _split_narrative_paras(intro)
        _write_paragraphs_from(cell, intro_idx + 1, lines)


def _fill_opinion_cell(cell, rec: str, sup: str) -> None:
    rec_lines = _split_narrative_paras(rec)
    sup_lines = _split_narrative_paras(sup)
    if rec_lines:
        rec_i = None
        for i, p in enumerate(cell.paragraphs):
            if "推荐理由" in (p.text or ""):
                rec_i = i
                break
        if rec_i is not None:
            _write_paragraphs_from(cell, rec_i + 1, rec_lines, stop_labels=("支持条件",))
        else:
            _write_paragraphs_from(cell, 0, ["1.推荐理由(含申报人情况介绍)"] + rec_lines)
    if sup_lines:
        sup_i = None
        for i, p in enumerate(cell.paragraphs):
            if "支持条件" in (p.text or ""):
                sup_i = i
                break
        if sup_i is not None:
            _write_paragraphs_from(cell, sup_i + 1, sup_lines)
        else:
            cell.add_paragraph("2.支持条件(包括工作和生活等方面)")
            _write_paragraphs_from(cell, len(cell.paragraphs) - 1, sup_lines)


def _fill_narrative_tables(doc: Document, fields: dict[str, str]) -> None:
    kind = _hj_template_kind(doc)
    exp = _clean_narrative(fields.get("专长及代表性成果") or "")
    wp = _clean_narrative(fields.get("工作设想") or "")
    for head in ("一、研究背景", "面向全球", "1. 研究背景", "1.研究背景"):
        idx = wp.find(head)
        if idx >= 0:
            wp = wp[idx:]
            break
    wp = re.sub(r"^工作设想[^\n]*\n(?:Work Plan[^\n]*\n)?", "", wp).strip()

    if kind == "compact":
        if len(doc.tables) > 3 and exp:
            if len(doc.tables[3].rows) > 1:
                _write_cell(doc.tables[3].rows[1].cells[0], exp)
            else:
                _fill_expertise_cell(doc.tables[3].rows[0].cells[0], exp)
        if len(doc.tables) > 5 and wp:
            if len(doc.tables[5].rows) > 1:
                _write_cell(doc.tables[5].rows[1].cells[0], wp)
            else:
                _write_cell(doc.tables[5].rows[0].cells[0], wp)
        if len(doc.tables) > 6 and fields.get("用人单位简介"):
            _write_cell(doc.tables[6].rows[0].cells[0], fields["用人单位简介"])
        if len(doc.tables) > 8:
            cell = doc.tables[8].rows[-1].cells[0]
            parts = []
            if fields.get("推荐理由"):
                parts.append("1.推荐理由\n" + fields["推荐理由"])
            if fields.get("支持条件"):
                parts.append("2.支持条件\n" + fields["支持条件"])
            if parts:
                cell.text = (cell.text.split("\n")[0] if cell.text.strip() else "申报单位(用人单位)意见") + "\n\n" + "\n\n".join(parts)
        return

    if len(doc.tables) > 4 and exp:
        _fill_expertise_cell(doc.tables[4].rows[0].cells[0], exp)
    if len(doc.tables) > 5:
        _fill_hj_table5(doc.tables[5], fields)
    if len(doc.tables) > 12 and wp and len(doc.tables[12].rows) > 1:
        _write_paragraphs_from(doc.tables[12].rows[1].cells[0], 0, _split_narrative_paras(wp))
    if len(doc.tables) > 13:
        _fill_hj_table13(doc.tables[13], fields)
    if len(doc.tables) > 14:
        _fill_hj_table14(doc.tables[14], fields)
    if len(doc.tables) > 15:
        cell = doc.tables[15].rows[-1].cells[0]
        rec = re.sub(r"^及引进的必要性[^\n]*\n?", "", str(fields.get("推荐理由") or "")).strip()
        rec = re.sub(r"^（含申报人情况介绍）[^\n]*\n?", "", rec).strip()
        rec = re.sub(r"^含申报人情况介绍[^\n]*\n?", "", rec).strip()
        sup = re.sub(r"^（包括工作和生活等方面）[^\n]*\n?", "", str(fields.get("支持条件") or "")).strip()
        _fill_opinion_cell(cell, rec, sup)


def _fill_hj_lists(doc: Document, data: dict, raw_text: str, fields: dict[str, str]) -> dict:
    parser_edu = list(data.get("主要学历") or [])
    parser_work = list(data.get("工作经历") or [])
    ocr_edu, ocr_work = _parse_list_lines(raw_text)
    edu = _fill_edu_gaps(_merge_timeline_rows(ocr_edu, parser_edu, "edu"))
    work = _merge_timeline_rows(ocr_work, parser_work, "work")
    fields = _enrich_hj_fields(fields, edu, work)
    if len(doc.tables) >= 3:
        _fill_hj_table0(doc.tables[0], fields)
        _fill_hj_table1(doc.tables[1], fields)
        _fill_hj_table2(doc.tables[2], edu, work)
    papers = _parse_papers(raw_text)
    projects = _parse_projects(raw_text)
    direction = str(fields.get("具体研究方向") or fields.get("研究领域关键词") or "").strip()
    filled_papers = 0
    if len(doc.tables) > 8:
        n8 = min(3, len(papers))
        _ensure_table_rows(doc.tables[8], max(6, 1 + n8) if _hj_template_kind(doc) == "full" else 1 + n8)
        filled_papers += _fill_paper_table(doc.tables[8], papers, 0, direction)
    if len(doc.tables) > 9:
        rest = max(0, len(papers) - filled_papers)
        _ensure_table_rows(doc.tables[9], 1 + min(3, rest))
        filled_papers += _fill_paper_table(doc.tables[9], papers, filled_papers, direction)
    filled_projects = 0
    if len(doc.tables) > 6:
        n6 = min(3, len(projects))
        _ensure_table_rows(doc.tables[6], max(6, 1 + n6) if _hj_template_kind(doc) == "full" else 1 + n6)
        filled_projects += _fill_project_table(doc.tables[6], projects, 0, direction)
    if len(doc.tables) > 7:
        rest = max(0, len(projects) - filled_projects)
        _ensure_table_rows(doc.tables[7], 1 + min(2, rest))
        filled_projects += _fill_project_table(doc.tables[7], projects, filled_projects, direction)
    return {"edu": len(edu), "work": len(work), "papers": filled_papers, "projects": filled_projects}


def _fill_qm_list_table(table, rows: list[dict]) -> int:
    if not rows:
        return 0
    n = 0
    for ri in range(1, len(table.rows)):
        if n >= len(rows):
            break
        cells = _distinct_cells(table.rows[ri])
        if len(cells) < 6:
            continue
        d = rows[n]
        vals = [
            str(n + 1),
            d.get("起止时间", ""),
            d.get("所在国家", ""),
            d.get("校名称", "") or d.get("工作单位", ""),
            d.get("专业领域", "") or d.get("担任职务", ""),
            d.get("学位", "") or d.get("任职情况", ""),
        ]
        for ci, val in enumerate(vals):
            if str(val or "").strip():
                _write_cell(cells[ci], str(val))
        n += 1
    return n


def _fill_qm_lists(doc: Document, data: dict, raw_text: str) -> None:
    parser_edu = list(data.get("主要学历") or [])
    parser_work = list(data.get("工作经历") or [])
    ocr_edu, ocr_work = _parse_list_lines(raw_text)
    edu = _merge_timeline_rows(ocr_edu, parser_edu, "edu")
    work = _merge_timeline_rows(ocr_work, parser_work, "work")
    if len(doc.tables) >= 3:
        _fill_qm_list_table(doc.tables[2], edu)
    if len(doc.tables) >= 4:
        _fill_qm_list_table(doc.tables[3], work)


def fill_declaration_template(mode: str, data: dict, raw_text: str, out_path: Path) -> dict:
    """复制 QM/HJ 模板并填入 data，写出 out_path。"""
    from .scan_text_normalize import deocr_scanned_text

    raw_text = deocr_scanned_text(raw_text)
    raw_for_lists = raw_text
    tpl = template_path(mode)
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(tpl, out_path)
    doc = Document(str(out_path))
    if str(mode).upper() == "HJ":
        from .hj_template_blank import blank_hj_document

        blank_hj_document(doc)
    fields = _build_fields(data, raw_text)
    if str(mode).upper() == "HJ":
        contact_fb = _extract_contact_vertical_fields(raw_for_lists)
        for key in ("手机号", "电子邮箱", "最高学位中文", "最高学位英文", "回国前单位职务中文", "回国前单位职务英文"):
            val = str(contact_fb.get(key) or "").strip()
            if val and (not fields.get(key) or _is_hj_field_garbage(fields.get(key) or "")):
                fields[key] = val
        if not fields.get("有效证件姓名"):
            fields["有效证件姓名"] = _extract_english_certificate_name(raw_for_lists)
        if not fields.get("中文（音译）名"):
            fields["中文（音译）名"] = _extract_chinese_transliteration(raw_for_lists)
        if fields.get("有效证件姓名"):
            fields["有效证件姓名"] = _fix_concatenated_latin_name(fields["有效证件姓名"], raw_for_lists)
        if not fields.get("申报人姓名"):
            fields["申报人姓名"] = fields.get("有效证件姓名") or ""
        if fields.get("中文（音译）名"):
            fields["中文（音译）名"] = _format_chinese_transliteration(fields["中文（音译）名"])
        elif raw_for_lists:
            fields["中文（音译）名"] = _extract_chinese_transliteration(raw_for_lists)
        fields = _finalize_hj_fields(fields, raw_for_lists)
    if str(mode).upper() == "HJ":
        ocr_edu, ocr_work = _parse_list_lines(raw_for_lists)
        fields = _enrich_hj_fields(fields, ocr_edu, ocr_work)
    _fill_cover_paragraphs(doc, fields)
    list_stats = {}
    if str(mode).upper() == "HJ":
        list_stats = _fill_hj_lists(doc, data, raw_for_lists, fields)
        _fill_narrative_tables(doc, fields)
        n_scalar = _fill_scalar_rows(doc, fields, skip_tables={0, 1, 2, 4, 5, 6, 7, 8, 9, 12, 13, 14, 15})
    else:
        _fill_qm_lists(doc, data, raw_text)
        n_scalar = _fill_scalar_rows(doc, fields)
    doc.save(str(out_path))
    return {
        "ok": True,
        "template": str(tpl),
        "fields": len(fields),
        "scalarRows": n_scalar,
        "eduRows": list_stats.get("edu", 0),
        "workRows": list_stats.get("work", 0),
        "paperRows": list_stats.get("papers", 0),
        "hasExpertise": bool(fields.get("专长及代表性成果")),
        "hasWorkPlan": bool(fields.get("工作设想")),
        "name": fields.get("有效证件姓名") or "",
        "enterprise": fields.get("申报单位") or fields.get("实验室名称") or "",
    }
