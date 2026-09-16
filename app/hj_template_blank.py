# -*- coding: utf-8 -*-
"""清除 HJ.docx 模板中的样例数据，保留版式与栏位标签。"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from .template_fill import _distinct_cells, _edu_work_row_bounds, _looks_like_label, _set_paragraph_text, _write_cell

_PICTURE_LOCALS = frozenset({"drawing", "pict", "object"})


def _strip_cell_pictures(cell) -> None:
    """去掉单元格内的图片/对象，空签字位不保留模板样例签名。"""
    tc = getattr(cell, "_tc", None)
    if tc is None:
        return
    for el in list(tc.iter()):
        if el.tag.rsplit("}", 1)[-1] not in _PICTURE_LOCALS:
            continue
        parent = el.getparent()
        if parent is not None:
            parent.remove(el)


def _drop_unused_images(doc: Document) -> None:
    used = set()
    embed_attr = qn("r:embed")
    for el in doc.element.iter():
        if el.tag.rsplit("}", 1)[-1] != "blip":
            continue
        rid = el.get(embed_attr)
        if rid:
            used.add(rid)
    for rel_id, rel in list(doc.part.rels.items()):
        if "image" not in str(rel.reltype).lower():
            continue
        if rel_id not in used:
            del doc.part.rels[rel_id]


def _reset_checkboxes(text: str) -> str:
    return str(text or "").replace("☑", "□").replace("☒", "□")


def _reset_cell_checkboxes(cell) -> None:
    for p in cell.paragraphs:
        t = str(p.text or "")
        if "☑" in t or "☒" in t:
            _set_paragraph_text(p, _reset_checkboxes(t))


def _clear_cover_paragraphs(doc: Document) -> None:
    rules = (
        (re.compile(r"(申报人姓名\s*\([^)]*\)\s*[:：])\s*.+$", re.I), r"\1            "),
        (re.compile(r"(申报单位[^:：\n]*[:：])\s*.+$", re.I), r"\1           "),
        (re.compile(r"(实验室名称[^:：\n]*[:：])\s*.+$", re.I), r"\1              "),
        (re.compile(r"(联\s*系\s*人[^:：\n]*[:：])\s*.+$", re.I), r"\1"),
        (re.compile(r"(联系人电话[^:：\n]*[:：])\s*.+$", re.I), r"\1"),
        (re.compile(r"(填表日期\s*\([^)]*\)\s*).*$", re.I), r"\1      年   月   日"),
    )
    for p in doc.paragraphs:
        text = str(p.text or "")
        if not text.strip():
            continue
        new_text = _reset_checkboxes(text)
        for pat, repl in rules:
            if pat.search(new_text):
                new_text = pat.sub(repl, new_text, count=1)
                break
        if re.search(r"所属二级学科及代码", new_text, re.I):
            continue
        if "□" in new_text or "☑" in text:
            new_text = _reset_checkboxes(new_text)
        if re.fullmatch(r"\s+", new_text) or new_text.strip() in {"", " "}:
            _set_paragraph_text(p, "")
        elif new_text != text:
            _set_paragraph_text(p, new_text)

    for i, p in enumerate(doc.paragraphs):
        t = str(p.text or "")
        if "二级学科" in t and "代码" in t:
            for j in range(i + 1, min(i + 4, len(doc.paragraphs))):
                nxt = doc.paragraphs[j]
                if not str(nxt.text or "").strip() or str(nxt.text or "").strip().isspace():
                    _set_paragraph_text(nxt, "")
                elif "前沿领域" in nxt.text or "关键核心" in nxt.text:
                    break
                else:
                    _set_paragraph_text(nxt, "")


def _is_data_value(text: str) -> bool:
    t = str(text or "").strip()
    if not t:
        return False
    if _looks_like_label(t):
        return False
    if re.search(r"[□☑☐]", t):
        return False
    if re.search(r"\d{4}[./-]\d", t):
        return True
    if re.search(r"@|\.com|\.cn|\.jp|\.edu", t, re.I):
        return True
    if re.match(r"^[A-Z]{1,3}\d{5,}$", t):
        return True
    if re.match(r"^[+\d][\d\-\s]{6,}$", t):
        return True
    if re.match(r"^[A-Za-z][A-Za-z .·'-]{1,60}$", t):
        return True
    if len(t) <= 40 and re.search(r"[\u4e00-\u9fff]", t) and not re.search(r"证件|号码|类型|国籍|性别|出生", t):
        return True
    if len(t) > 40:
        return True
    return False


def _clear_identity_tables(doc: Document) -> None:
    if len(doc.tables) < 2:
        return
    t0 = doc.tables[0]
    for row in t0.rows:
        cells = _distinct_cells(row)
        row_text = " ".join(c.text for c in cells)
        for i, cell in enumerate(cells):
            txt = str(cell.text or "").strip()
            if not txt:
                continue
            if "□" in txt or "☑" in txt:
                _reset_cell_checkboxes(cell)
                continue
            if i >= 2 and not _looks_like_label(txt):
                _write_cell(cell, "")
            elif i > 0 and _is_data_value(txt) and not _looks_like_label(txt):
                _write_cell(cell, "")
        if "性别" in row_text and "出生" in row_text:
            if len(cells) >= 2:
                _write_cell(cells[1], "")
            if len(cells) >= 4:
                _write_cell(cells[3], "")
            if len(cells) >= 6:
                _write_cell(cells[5], "")

    t1 = doc.tables[1]
    for ri, row in enumerate(t1.rows):
        cells = _distinct_cells(row)
        for i, cell in enumerate(cells):
            txt = str(cell.text or "").strip()
            if not txt:
                continue
            if "□" in txt or "☑" in txt:
                _reset_cell_checkboxes(cell)
                continue
            if ri in (0, 1, 2, 3, 4, 5) and i >= len(cells) - 1 and _is_data_value(txt):
                _write_cell(cell, "")
            elif ri == 8 and i in (1, 3, 5) and _is_data_value(txt):
                _write_cell(cell, "")
            elif ri in (9, 10, 11) and i >= 1 and _is_data_value(txt):
                _write_cell(cell, "")


def _clear_list_table(table, data_cols: tuple[int, ...]) -> None:
    for ri in range(1, len(table.rows)):
        cells = _distinct_cells(table.rows[ri])
        if len(cells) < 2:
            continue
        for ci in data_cols:
            if ci < len(cells):
                _write_cell(cells[ci], "")


def _clear_table2(doc: Document) -> None:
    if len(doc.tables) < 3:
        return
    table = doc.tables[2]
    edu_start, edu_end, work_start = _edu_work_row_bounds(table)
    for ri in range(edu_start, edu_end):
        cells = _distinct_cells(table.rows[ri])
        for ci in range(1, min(6, len(cells))):
            _write_cell(cells[ci], "")
    for ri in range(work_start, len(table.rows)):
        cells = _distinct_cells(table.rows[ri])
        if any("时间" in c.text and "Time" in c.text for c in cells[1:3] if len(cells) > 2):
            continue
        for ci in range(1, min(6, len(cells))):
            _write_cell(cells[ci], "")


def _trim_narrative_cell(cell, keep_lines: int) -> None:
    paras = list(cell.paragraphs)
    if not paras:
        return
    kept = 0
    for p in paras:
        t = str(p.text or "")
        if not t.strip():
            continue
        kept += 1
        if "☑" in t or "☒" in t:
            _set_paragraph_text(p, _reset_checkboxes(t))
        if kept > keep_lines:
            _set_paragraph_text(p, "")


def _clear_narrative_tables(doc: Document) -> None:
    n = len(doc.tables)
    if n > 4:
        _trim_narrative_cell(doc.tables[4].rows[0].cells[0], 3)
    if n > 5:
        _trim_narrative_cell(doc.tables[5].rows[0].cells[0], 8)
    for ti in range(6, min(11, n)):
        _clear_list_table(doc.tables[ti], tuple(range(1, 6)))
    if n > 11:
        cell = doc.tables[11].rows[0].cells[0]
        kept = False
        for p in cell.paragraphs:
            t = str(p.text or "").strip()
            if not t:
                continue
            if not kept:
                kept = True
                continue
            _set_paragraph_text(p, "")
    if n > 12 and len(doc.tables[12].rows) > 1:
        _write_cell(doc.tables[12].rows[1].cells[0], "")
    if n > 13:
        for ri in range(1, len(doc.tables[13].rows)):
            cell = doc.tables[13].rows[ri].cells[0]
            _strip_cell_pictures(cell)
            for p in cell.paragraphs:
                txt = str(p.text or "")
                if "☑" in txt or "☒" in txt:
                    _set_paragraph_text(p, _reset_checkboxes(txt))
    if n > 14:
        cell = doc.tables[14].rows[0].cells[0]
        past_intro = False
        for p in cell.paragraphs:
            t = str(p.text or "")
            if "☑" in t or "☒" in t:
                _set_paragraph_text(p, _reset_checkboxes(t))
            if past_intro:
                _set_paragraph_text(p, "")
            elif "简介" in t and ("300" in t or "实验室建设" in t):
                past_intro = True
            elif "上级部门" in t or "按隶属关系" in t:
                nt = re.sub(r"([:：])\s*.+$", r"\1 ", t)
                if nt != t:
                    _set_paragraph_text(p, nt)
    if n > 15:
        cell = doc.tables[15].rows[-1].cells[0]
        _strip_cell_pictures(cell)
        for p in cell.paragraphs:
            t = str(p.text or "")
            if "推荐理由" in t or "支持条件" in t:
                continue
            if t.strip():
                _set_paragraph_text(p, "")


def blank_hj_document(doc: Document) -> None:
    """清除已打开的 HJ 模板文档中的样例填写内容。"""
    _clear_cover_paragraphs(doc)
    _clear_identity_tables(doc)
    if len(doc.tables) > 3:
        _clear_table2(doc)
        _clear_list_table(doc.tables[3], (1, 2, 3))
    _clear_narrative_tables(doc)
    _drop_unused_images(doc)


def make_blank_template(src: Path, out: Path) -> Path:
    import shutil

    if not src.is_file():
        raise FileNotFoundError("源模板不存在：" + str(src))
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(src, out)
    doc = Document(str(out))
    blank_hj_document(doc)
    doc.save(str(out))
    return out
