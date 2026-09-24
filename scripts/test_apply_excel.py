# -*- coding: utf-8 -*-
"""Excel 跨列替换：制表符 replace 应按列写回，避免整段挤进首格。"""
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))

from apply_excel import SEP_ROW, _apply_joined, apply_edits_to_sheets


def _row(*vals):
    return {"r": 1, "cells": [{"c": i + 1, "text": v, "dirty": False, "orig": v} for i, v in enumerate(vals)]}


def test_tabbed_replace_splits_columns():
    row = _row("3", "2024.07.01", "至今", "比利时", "终身副教授", "Lifelong Associate Professor", "副教授", "根特大学")
    joined = SEP_ROW.join(c["text"] for c in row["cells"])
    find = "3" + SEP_ROW + "2024.07.01" + SEP_ROW + "至今" + SEP_ROW + "比利时" + SEP_ROW + "终身副教授"
    rep = "3" + SEP_ROW + "2024.07.01" + SEP_ROW + "至今" + SEP_ROW + "比利时" + SEP_ROW + "副教授"
    p = joined.find(find)
    assert p >= 0, joined
    ok = _apply_joined(row["cells"], joined, p, p + len(find), rep)
    assert ok
    assert row["cells"][0]["text"] == "3"
    assert row["cells"][1]["text"] == "2024.07.01"
    assert row["cells"][2]["text"] == "至今"
    assert row["cells"][3]["text"] == "比利时"
    assert row["cells"][4]["text"] == "副教授"
    assert row["cells"][5]["text"] == "Lifelong Associate Professor"
    print("tabbed replace ok")


def test_hit_rows_integration():
    sheets = [{"name": "s1", "rows": [_row("3", "2024.07.01", "至今", "比利时", "终身副教授", "Lifelong")]}]
    find = "3" + SEP_ROW + "2024.07.01" + SEP_ROW + "至今" + SEP_ROW + "比利时" + SEP_ROW + "终身副教授"
    rep = "3" + SEP_ROW + "2024.07.01" + SEP_ROW + "至今" + SEP_ROW + "比利时" + SEP_ROW + "副教授"
    res = apply_edits_to_sheets(sheets, [{"find": find, "replace": rep}])
    assert res[0]["status"] == "hit-span"
    cells = sheets[0]["rows"][0]["cells"]
    assert cells[4]["text"] == "副教授"
    assert cells[1]["text"] == "2024.07.01"
    print("hit-rows ok", res)


def test_keyword_label_fills_empty_cell_below():
    sheets = [{
        "name": "s1",
        "rows": [
            {"r": 96, "cells": [{"c": 2, "text": "研究领域关键词（不超过5个）", "dirty": False, "orig": "研究领域关键词（不超过5个）"}]},
            {"r": 97, "cells": [{"c": 2, "text": "Key words (no more than 5 )", "dirty": False, "orig": "Key words (no more than 5 )"}]},
            {"r": 105, "cells": [{"c": 2, "text": "2.代表性科研项目", "dirty": False, "orig": "2.代表性科研项目"}]},
        ],
    }]
    kws = "计算机体系结构、能效计算、非易失性处理器、存内计算、系统安全"
    res = apply_edits_to_sheets(sheets, [{"find": "研究领域关键词（不超过5个）", "replace": kws}])
    assert res[0]["status"].startswith("hit"), res
    rows = {row["r"]: row["cells"][0]["text"] for row in sheets[0]["rows"]}
    assert rows[96] == "研究领域关键词（不超过5个）", rows[96]
    assert rows[97].startswith("Key words"), rows[97]
    assert kws in rows[98], rows
    assert rows[105].startswith("2.代表性科研项目")
    print("keyword fill-below ok")


def test_keyword_write_uses_merged_value_cell():
    from openpyxl import Workbook
    from apply_excel import load_openpyxl_sheets, apply_edits_to_sheets, write_openpyxl

    src = ROOT / "scripts" / "_tmp_kw.xlsx"
    out = ROOT / "scripts" / "_tmp_kw_out.xlsx"
    wb = Workbook()
    ws = wb.active
    ws["B1"] = "研究领域关键词（不超过5个）"
    ws["B2"] = "Key words (no more than 5 )"
    ws.merge_cells("B2:E2")
    ws.merge_cells("B3:E8")
    ws["B9"] = "2.代表性科研项目(主持/参与)"
    wb.save(src)
    wb.close()
    book, sheets = load_openpyxl_sheets(src)
    kws = "计算机体系结构、能效计算"
    res = apply_edits_to_sheets(sheets, [{"find": "研究领域关键词（不超过5个）", "replace": kws}])
    assert res[0]["status"].startswith("hit"), res
    write_openpyxl(book, sheets, out)
    from openpyxl import load_workbook
    wb2 = load_workbook(out)
    ws2 = wb2.active
    assert ws2["B1"].value == "研究领域关键词（不超过5个）"
    assert "Key words" in str(ws2["B2"].value)
    assert kws in str(ws2["B3"].value or "")
    assert ws2["B9"].value.startswith("2.代表性科研项目")
    wb2.close()
    for p in (src, out):
        if p.exists():
            p.unlink()
    print("keyword merged write ok")


def test_merged_slave_clear_does_not_crash():
    from openpyxl import Workbook
    from apply_excel import load_openpyxl_sheets, apply_edits_to_sheets, write_openpyxl

    src = ROOT / "scripts" / "_tmp_merge.xlsx"
    out = ROOT / "scripts" / "_tmp_merge_out.xlsx"
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "学术成就：原文"
    ws.merge_cells("A1:C1")
    ws["A2"] = "保持不变"
    wb.save(src)
    wb.close()
    book, sheets = load_openpyxl_sheets(src)
    find = "学术成就：原文"
    res = apply_edits_to_sheets(sheets, [{"find": find, "replace": "学术成就：已改"}])
    assert res[0]["status"].startswith("hit")
    # 合并从格被跨列改写标成清空时，不能对 MergedCell 赋值
    row = sheets[0]["rows"][0]
    row["cells"].append({"c": 2, "text": "", "dirty": True, "orig": None})
    row["cells"].append({"c": 3, "text": "", "dirty": True, "orig": None})
    write_openpyxl(book, sheets, out)
    book2, sheets2 = load_openpyxl_sheets(out)
    texts = [c["text"] for row in sheets2[0]["rows"] for c in row["cells"]]
    assert any("已改" in t for t in texts), texts
    for p in (src, out):
        if p.exists():
            p.unlink()
    print("merged cell write ok")


if __name__ == "__main__":
    test_tabbed_replace_splits_columns()
    test_hit_rows_integration()
    test_keyword_label_fills_empty_cell_below()
    test_keyword_write_uses_merged_value_cell()
    test_merged_slave_clear_does_not_crash()
    print("ALL OK")
